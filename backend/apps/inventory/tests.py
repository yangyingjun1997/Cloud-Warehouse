from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.apps import apps
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.test import override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from PIL import Image
from rest_framework.test import APIClient
import shutil
import tempfile
from pathlib import Path
from io import BytesIO
from importlib import import_module

from .check_services import (
    create_check_task,
    correct_check_line,
    reopen_check_task,
    review_check_task,
    scan_check_value,
    start_check_task,
    submit_check_task,
)
from .conversion_services import (
    InventoryConversionError,
    convert_assets_to_stock,
    convert_stock_to_assets,
)
from .import_services import (
    apply_inventory_import,
    create_import_batch,
    parse_inventory_workbook,
    build_inventory_import_template,
    validate_inventory_records,
)
from .lifecycle_services import (
    archive_asset_attachment,
    create_handover,
    create_maintenance_record,
    finish_handover,
    finish_maintenance_record,
)
from .models import (
    Asset,
    AssetAttachment,
    AssetHandover,
    AssetLifecycleEvent,
    AssetNetworkEndpoint,
    Category,
    ComponentLink,
    CompositeUnit,
    ImportBatch,
    InventoryConversion,
    InventoryCheckAdjustment,
    InventoryCheckLine,
    InventoryCheckTask,
    ItemType,
    Location,
    MaintenanceRecord,
    StockItem,
    Warehouse,
)
from .search import asset_search_query
from apps.common.models import OperationAuditLog
from apps.notifications.models import Notification
from apps.workflow.models import InventoryReservation, WorkflowRequest, WorkflowRequestLine
from apps.workflow.services import process_request

User = get_user_model()
TEST_MEDIA_ROOT = Path(tempfile.mkdtemp(prefix='warehouse-test-media-'))


class AssetStatusDimensionTests(TestCase):
    def setUp(self):
        category = Category.objects.create(name='状态测试分类', code='STATUS-TEST')
        self.item_type = ItemType.objects.create(name='状态测试类型', code='STATUS-TYPE', category=category)
        self.warehouse = Warehouse.objects.create(name='状态测试仓', code='STATUS-WH')

    def test_legacy_status_is_mapped_when_asset_is_created(self):
        asset = Asset.objects.create(
            name='借出设备', item_type=self.item_type, warehouse=self.warehouse,
            status=Asset.Status.BORROWED,
        )

        self.assertEqual(asset.location_state, Asset.LocationState.OUT_ON_LOAN)
        self.assertEqual(asset.availability_state, Asset.AvailabilityState.UNAVAILABLE)
        self.assertEqual(asset.quality_state, Asset.QualityState.NORMAL)
        self.assertEqual(asset.disposition_state, Asset.DispositionState.INTERNAL)

    def test_status_mapping_covers_every_legacy_status(self):
        status_values = {value for value, _ in Asset.Status.choices}
        self.assertEqual(status_values, set(Asset.LEGACY_STATUS_DIMENSIONS))
        for dimensions in Asset.LEGACY_STATUS_DIMENSIONS.values():
            self.assertEqual(
                set(dimensions),
                {'location_state', 'availability_state', 'quality_state', 'disposition_state'},
            )

    def test_status_update_fields_also_persists_dimensions(self):
        asset = Asset.objects.create(
            name='待维修设备', item_type=self.item_type, warehouse=self.warehouse,
        )
        asset.status = Asset.Status.REPAIRING
        asset.save(update_fields=['status', 'updated_at'])

        asset.refresh_from_db()
        self.assertEqual(asset.status, Asset.Status.REPAIRING)
        self.assertEqual(asset.location_state, Asset.LocationState.EXTERNAL_REPAIR)
        self.assertEqual(asset.availability_state, Asset.AvailabilityState.UNAVAILABLE)
        self.assertEqual(asset.quality_state, Asset.QualityState.REPAIRING)

    def test_non_status_edit_preserves_reservation_availability(self):
        asset = Asset.objects.create(
            name='预约设备', item_type=self.item_type, warehouse=self.warehouse,
        )
        Asset.objects.filter(pk=asset.pk).update(
            availability_state=Asset.AvailabilityState.RESERVED,
        )
        asset.remarks = '补充资料不应释放预约'
        asset.save()

        asset.refresh_from_db()
        self.assertEqual(asset.availability_state, Asset.AvailabilityState.RESERVED)

    def test_assembled_component_is_not_individually_available(self):
        asset = Asset.objects.create(
            name='已组装组件', item_type=self.item_type, warehouse=self.warehouse,
            status=Asset.Status.ASSEMBLED,
        )
        self.assertEqual(asset.availability_state, Asset.AvailabilityState.UNAVAILABLE)

    def test_historical_backfill_restores_dimensions_from_legacy_status(self):
        asset = Asset.objects.create(
            name='历史借出设备', item_type=self.item_type, warehouse=self.warehouse,
            status=Asset.Status.BORROWED,
        )
        Asset.objects.filter(pk=asset.pk).update(
            location_state='', availability_state='', quality_state='', disposition_state='',
        )

        migration = import_module('apps.inventory.migrations.0021_asset_status_dimensions')
        migration.backfill_asset_status_dimensions(apps, None)

        asset.refresh_from_db()
        self.assertEqual(asset.location_state, Asset.LocationState.OUT_ON_LOAN)
        self.assertEqual(asset.availability_state, Asset.AvailabilityState.UNAVAILABLE)
        self.assertEqual(asset.quality_state, Asset.QualityState.NORMAL)
        self.assertEqual(asset.disposition_state, Asset.DispositionState.INTERNAL)


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class AssetLifecycleTests(TestCase):
    def setUp(self):
        self.employee = User.objects.create_user(username='lifecycle-employee', password='password123')
        self.operator = User.objects.create_user(username='lifecycle-operator', password='password123')
        self.operator.groups.add(Group.objects.create(name='warehouse_staff'))
        category = Category.objects.create(name='机器人资产', code='ROBOT-ASSET')
        item_type = ItemType.objects.create(name='机器狗', code='ROBOT-DOG', category=category)
        warehouse = Warehouse.objects.create(name='生命周期测试仓', code='LIFE-WH')
        self.asset = Asset.objects.create(
            name='测试机器人本体',
            item_type=item_type,
            warehouse=warehouse,
            current_holder=self.employee,
            status=Asset.Status.BORROWED,
        )

    def test_manual_maintenance_is_audited_and_completion_restores_previous_status(self):
        record = create_maintenance_record(
            asset=self.asset,
            actor=self.operator,
            repair_type=MaintenanceRecord.RepairType.EXTERNAL,
            fault_description='机器人无法开机',
            service_provider='测试维修商',
        )
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.REPAIRING)
        self.assertEqual(record.previous_asset_status, Asset.Status.BORROWED)
        self.assertTrue(OperationAuditLog.objects.filter(action='asset_maintenance_created', target_id=record.id).exists())

        finish_maintenance_record(record, self.operator, resolution='更换主板', actual_cost='1200.00')
        self.asset.refresh_from_db()
        record.refresh_from_db()
        self.assertEqual(record.status, MaintenanceRecord.Status.COMPLETED)
        self.assertEqual(self.asset.status, Asset.Status.BORROWED)
        self.assertEqual(record.actual_cost, 1200)
        self.assertGreaterEqual(self.asset.lifecycle_events.count(), 2)

    def test_processed_repair_request_automatically_creates_maintenance_record(self):
        request_obj = WorkflowRequest.objects.create(
            request_type=WorkflowRequest.RequestType.REPAIR,
            applicant=self.employee,
            status=WorkflowRequest.Status.WAITING_WAREHOUSE,
            reason='关节电机异常',
        )
        WorkflowRequestLine.objects.create(request=request_obj, asset=self.asset, quantity=1)

        process_request(request_obj, self.operator)

        record = MaintenanceRecord.objects.get(asset=self.asset, source_request=request_obj)
        self.asset.refresh_from_db()
        self.assertEqual(record.fault_description, '关节电机异常')
        self.assertEqual(record.previous_asset_status, Asset.Status.BORROWED)
        self.assertEqual(self.asset.status, Asset.Status.REPAIRING)

    def test_handover_confirmation_does_not_change_inventory_status_or_holder(self):
        handover = create_handover(
            asset=self.asset,
            actor=self.operator,
            from_holder_name='原使用人',
            to_holder_name='新使用人',
            handover_date=timezone.localdate(),
        )
        finish_handover(handover, self.operator)

        self.asset.refresh_from_db()
        handover.refresh_from_db()
        self.assertEqual(handover.status, AssetHandover.Status.CONFIRMED)
        self.assertEqual(self.asset.status, Asset.Status.BORROWED)
        self.assertEqual(self.asset.current_holder, self.employee)

    def test_attachment_archive_and_lifecycle_page_permissions(self):
        attachment = archive_asset_attachment(
            asset=self.asset,
            actor=self.operator,
            file=SimpleUploadedFile('repair-report.pdf', b'%PDF-test', content_type='application/pdf'),
            display_name='维修检测报告',
            category=AssetAttachment.Category.REPAIR_REPORT,
        )
        self.assertTrue(attachment.file.name.endswith('repair-report.pdf'))
        self.assertTrue(AssetLifecycleEvent.objects.filter(
            asset=self.asset,
            related_object_id=str(attachment.id),
        ).exists())

        self.client.force_login(self.operator)
        response = self.client.get(f'/warehouse/assets/{self.asset.id}/lifecycle/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '生命周期时间线')
        self.assertContains(response, '维修检测报告')
        response = self.client.get(f'/warehouse/attachments/{attachment.id}/download/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment;', response['Content-Disposition'])

        self.client.force_login(self.employee)
        response = self.client.get(f'/warehouse/assets/{self.asset.id}/lifecycle/')
        self.assertRedirects(response, '/')
        response = self.client.get(f'/warehouse/attachments/{attachment.id}/download/')
        self.assertEqual(response.status_code, 404)

    def test_scrap_button_submits_existing_workflow_approval(self):
        self.client.force_login(self.operator)
        response = self.client.post(f'/warehouse/assets/{self.asset.id}/lifecycle/', {
            'action': 'request_scrap',
            'scrap-reason': '核心结构损坏且无维修价值',
        })
        request_obj = WorkflowRequest.objects.get(
            request_type=WorkflowRequest.RequestType.SCRAP,
            lines__asset=self.asset,
        )
        self.assertRedirects(response, f'/workflow/{request_obj.id}/')
        self.assertEqual(request_obj.status, WorkflowRequest.Status.PENDING)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.BORROWED)


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class AssetApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.employee = User.objects.create_user(username='employee', password='password123')
        self.operator = User.objects.create_user(username='operator', password='password123', is_staff=True)
        self.operator.groups.add(Group.objects.create(name='warehouse_staff'))
        category = Category.objects.create(name='测试分类', code='TEST')
        item_type = ItemType.objects.create(name='测试类型', code='TEST-TYPE', category=category)
        self.asset = Asset.objects.create(
            name='测试资产',
            item_type=item_type,
            asset_code='TEST-ZC-001',
            serial_number='SN-001',
            qr_value='QR-001',
        )

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEST_MEDIA_ROOT, ignore_errors=True)

    def test_lookup_and_qr_code(self):
        self.client.force_authenticate(self.employee)
        response = self.client.get(f'/api/inventory/assets/lookup/?code={self.asset.asset_code}')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['id'], str(self.asset.id))
        self.assertEqual(response.data['system_asset_no'], self.asset.system_asset_no)

        response = self.client.get(f'/api/inventory/assets/lookup/?code={self.asset.system_asset_no}')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['id'], str(self.asset.id))

        self.asset.manufacturer_barcode = 'MFG-CODE-001'
        self.asset.save(update_fields=['manufacturer_barcode', 'updated_at'])
        for identifier in [self.asset.serial_number, self.asset.qr_value, self.asset.manufacturer_barcode]:
            response = self.client.get(f'/api/inventory/assets/lookup/?code={identifier}')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.data['id'], str(self.asset.id))

        response = self.client.get(f'/api/inventory/assets/{self.asset.id}/qr-code/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'image/png')
        self.assertTrue(response.content.startswith(b'\x89PNG'))

    def test_lookup_requires_authentication(self):
        response = self.client.get(f'/api/inventory/assets/lookup/?code={self.asset.asset_code}')
        self.assertEqual(response.status_code, 401)

    def test_operator_can_upload_valid_asset_photo(self):
        self.client.force_authenticate(self.operator)
        buffer = BytesIO()
        Image.new('RGB', (2, 2), color='white').save(buffer, format='PNG')
        photo = SimpleUploadedFile('asset.png', buffer.getvalue(), content_type='image/png')
        response = self.client.post(
            f'/api/inventory/assets/{self.asset.id}/upload-photo/',
            {'photo': photo},
            format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data['photo_url'])

    def test_operator_cannot_upload_non_image_as_asset_photo(self):
        self.client.force_authenticate(self.operator)
        photo = SimpleUploadedFile('asset.png', b'not-an-image', content_type='image/png')
        response = self.client.post(
            f'/api/inventory/assets/{self.asset.id}/upload-photo/',
            {'photo': photo},
            format='multipart',
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('有效的图片', response.data['detail'])

    def test_anonymous_api_list_requires_authentication_by_default(self):
        response = self.client.get('/api/inventory/network-endpoints/')

        self.assertEqual(response.status_code, 401)

    def test_network_endpoint_has_its_own_resource(self):
        endpoint = AssetNetworkEndpoint.objects.create(
            asset=self.asset,
            interface_name='LAN',
            ip_address='192.168.1.10',
            mac_address='00:11:22:33:44:55',
            is_primary=True,
        )
        self.client.force_authenticate(self.operator)
        response = self.client.get(f'/api/inventory/network-endpoints/{endpoint.id}/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(str(response.data['asset']), str(self.asset.id))
        self.assertEqual(response.data['ip_address'], '192.168.1.10')

        self.client.force_authenticate(self.employee)
        response = self.client.get(f'/api/inventory/network-endpoints/{endpoint.id}/')
        self.assertEqual(response.status_code, 403)

    def test_asset_identifiers_are_serialized(self):
        self.asset.manufacturer = '测试厂家'
        self.asset.model = '测试型号'
        self.asset.manufacturer_barcode = 'MFG-BARCODE-001'
        self.asset.purchase_order_no = 'PO-001'
        self.asset.save()

        self.client.force_authenticate(self.operator)
        response = self.client.get(f'/api/inventory/assets/{self.asset.id}/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['manufacturer'], '测试厂家')
        self.assertEqual(response.data['manufacturer_barcode'], 'MFG-BARCODE-001')
        self.assertEqual(response.data['purchase_order_no'], 'PO-001')
        self.assertTrue(response.data['can_request'])
        self.assertFalse(response.data['can_return'])

    def test_reserved_asset_is_not_requestable_in_api(self):
        request_obj = WorkflowRequest.objects.create(
            request_type=WorkflowRequest.RequestType.BORROW,
            applicant=self.employee,
        )
        line = WorkflowRequestLine.objects.create(request=request_obj, asset=self.asset, quantity=1)
        InventoryReservation.objects.create(request_line=line, asset=self.asset, quantity=1)

        self.client.force_authenticate(self.employee)
        response = self.client.get(f'/api/inventory/assets/{self.asset.id}/')

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['availability_state'], Asset.AvailabilityState.RESERVED)
        self.assertFalse(response.data['can_request'])

    def test_search_supports_chinese_pinyin_initials_aliases_and_remarks(self):
        self.asset.name = '机器人底盘'
        self.asset.search_aliases = '巡检底盘'
        self.asset.remarks = '展厅备用设备'
        self.asset.save()
        self.client.force_authenticate(self.employee)

        for query in ['jqr', '巡检', '展厅备用']:
            response = self.client.get('/api/inventory/assets/search/', {'q': query})
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data['count'], 1)
            self.assertEqual(response.data['results'][0]['id'], str(self.asset.id))

    def test_employee_stock_api_hides_exact_quantity_and_asset_sensitive_fields(self):
        StockItem.objects.create(name='测试耗材', code='CONSUMABLE-001', quantity=3)
        self.client.force_authenticate(self.employee)
        response = self.client.get('/api/inventory/stock-items/')
        self.assertEqual(response.status_code, 200, response.data)
        stock = response.data['results'][0]
        self.assertTrue(stock['can_request'])
        self.assertNotIn('quantity', stock)

        response = self.client.get(f'/api/inventory/assets/{self.asset.id}/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertNotIn('warehouse', response.data)
        self.assertNotIn('current_holder', response.data)

    def test_warehouse_staff_cannot_change_asset_state_or_stock_quantity_through_api(self):
        stock_item = StockItem.objects.create(name='测试耗材', code='CONSUMABLE-LOCKED', quantity=3)
        self.client.force_authenticate(self.operator)

        response = self.client.patch(
            f'/api/inventory/assets/{self.asset.id}/',
            {'status': Asset.Status.BORROWED},
            format='json',
        )
        self.assertEqual(response.status_code, 400, response.data)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.IN_STOCK)

        response = self.client.patch(
            f'/api/inventory/stock-items/{stock_item.id}/',
            {'quantity': 9},
            format='json',
        )
        self.assertEqual(response.status_code, 400, response.data)
        stock_item.refresh_from_db()
        self.assertEqual(stock_item.quantity, 3)

    def test_only_administrator_can_change_base_data_and_api_changes_are_audited(self):
        self.client.force_authenticate(self.operator)
        response = self.client.post('/api/inventory/categories/', {
            'name': '仓库人员不能创建',
            'code': 'STAFF-DENIED',
        }, format='json')
        self.assertEqual(response.status_code, 403, response.data)

        administrator = User.objects.create_superuser(username='api-admin', password='password123', email='')
        self.client.force_authenticate(administrator)
        response = self.client.post('/api/inventory/categories/', {
            'name': '接口创建分类',
            'code': 'API-CATEGORY',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue(OperationAuditLog.objects.filter(
            action='category_created', target_id=response.data['id'], actor=administrator,
        ).exists())


class InventoryConversionAndImportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(username='inventory-admin', password='password123', email='')
        self.category = Category.objects.create(name='转换测试分类', code='CONVERT-CAT')
        self.item_type = ItemType.objects.create(
            name='转换测试类型', code='CONVERT-TYPE', category=self.category,
            unit='件', safety_stock=2,
        )
        self.warehouse = Warehouse.objects.create(name='转换测试仓库', code='CONVERT-WH')
        self.location = Location.objects.create(
            warehouse=self.warehouse, name='转换测试库位', code='A-01',
        )
        self.stock = StockItem.objects.create(
            name='转换测试配件', code='CONVERT-STOCK', item_type=self.item_type,
            warehouse=self.warehouse, location=self.location, quantity=3, unit='件',
        )

    def _workbook(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '库存导入'
        sheet.append([
            '管理方式', '物品名称', '物品类型', '物品类型编码', '仓库', '库位',
            '系统编号', '厂家出厂SN序列号', '耗材配件编码', '库存数量', '单位',
            '最低库存提醒值', '状态', '备注',
        ])
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_download_template_uses_current_columns_and_five_examples(self):
        workbook = load_workbook(build_inventory_import_template(), data_only=True)
        sheet = workbook['库存导入']

        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            [
                '管理方式*', '物品名称*', '物品类型*', '物品类型编码', '仓库*', '库位',
                '仓库系统资产编号', '公司 ERP 资产编码', '生产厂家 SN / 序列号', '生产厂家条码 / 二维码内容', '本系统二维码内容',
                '耗材配件编码', '库存数量', '单位', '最低库存提醒值', '厂家', '型号',
                '供应商', '单价（元）', '状态', '备注',
            ],
        )
        records, issues, summary = parse_inventory_workbook(build_inventory_import_template().getvalue())
        self.assertEqual(len(records), 5)
        self.assertEqual(summary['valid_row_count'], 5)
        self.assertFalse([issue for issue in issues if issue['severity'] == 'error'])
        self.assertEqual(sheet['G4'].value, 'AST-000003')
        self.assertEqual(sheet['H4'].value, 'DZSB-D7-21-0001-600001')
        self.assertIn('填写说明', summary['sheets'])
        notes = workbook['填写说明']
        note_text = '\n'.join(str(row[1].value) for row in notes.iter_rows(min_row=2, max_col=2))
        self.assertIn('管理方式、物品名称、物品类型、仓库', note_text)
        self.assertIn('行政', note_text)
        self.assertIn('公司内部字段', note_text)

    def test_new_asset_gets_immutable_system_number_without_erp_code(self):
        asset = Asset.objects.create(
            name='未赋码资产', item_type=self.item_type,
            warehouse=self.warehouse, location=self.location,
        )
        self.assertIsNone(asset.asset_code)
        self.assertRegex(asset.system_asset_no, r'^AST-\d{6}$')
        self.assertEqual(asset.qr_value, asset.system_asset_no)

        asset.asset_code = 'DZSB-D7-21-0001-600001'
        asset.save()
        asset.refresh_from_db()
        self.assertEqual(asset.qr_value, asset.system_asset_no)

    def test_system_number_is_unique_and_searchable(self):
        asset = Asset.objects.create(
            name='编号检索资产', item_type=self.item_type,
            warehouse=self.warehouse, location=self.location,
        )
        self.assertEqual(Asset.objects.filter(asset_search_query(asset.system_asset_no)).count(), 1)

    def test_multiple_uncoded_assets_coexist(self):
        for index in range(3):
            Asset.objects.create(
                name=f'未赋码资产{index}', item_type=self.item_type,
                warehouse=self.warehouse, location=self.location,
            )
        self.assertEqual(
            Asset.objects.filter(asset_code__isnull=True, warehouse=self.warehouse).count(),
            3,
        )

    def test_controlled_conversion_preserves_history_and_balances_quantity(self):
        conversion, assets = convert_stock_to_assets(
            actor=self.admin, stock_item_id=self.stock.id, quantity=2,
            serial_numbers='CONVERT-SN-001\nCONVERT-SN-002', note='拆出两件独立追踪',
        )
        self.stock.refresh_from_db()
        self.assertEqual(self.stock.quantity, 1)
        self.assertEqual(len(assets), 2)
        self.assertEqual(conversion.direction, InventoryConversion.Direction.STOCK_TO_ASSET)
        self.assertTrue(OperationAuditLog.objects.filter(action='stock_converted_to_assets', target_id=conversion.id).exists())

        reverse_conversion, target = convert_assets_to_stock(
            actor=self.admin, asset_ids=[asset.id for asset in assets],
            target_stock_item_id=self.stock.id, note='恢复按数量管理',
        )
        target.refresh_from_db()
        self.assertEqual(target.quantity, 3)
        self.assertEqual(reverse_conversion.assets.count(), 2)
        self.assertEqual(
            set(Asset.objects.filter(pk__in=[asset.id for asset in assets]).values_list('status', flat=True)),
            {Asset.Status.CONVERTED_TO_STOCK},
        )

    def test_only_in_stock_assets_can_convert_to_stock(self):
        asset = Asset.objects.create(
            name='借出资产', item_type=self.item_type, warehouse=self.warehouse,
            location=self.location, status=Asset.Status.BORROWED,
        )
        with self.assertRaisesMessage(InventoryConversionError, '在库'):
            convert_assets_to_stock(actor=self.admin, asset_ids=[asset.id])

    def test_incremental_import_creates_and_updates_unique_records(self):
        content = self._workbook([
            ['单件资产', '新雷达', '雷达', 'IMPORT-LIDAR', self.warehouse.name, self.location.code, '', 'IMPORT-SN-001', '', 1, '件', 1, '在库', '增量导入测试'],
            ['耗材配件', self.stock.name, self.item_type.name, self.item_type.code, self.warehouse.name, self.location.code, '', '', self.stock.code, 9, '件', 3, '', '更新库存'],
        ])
        records, issues, summary = parse_inventory_workbook(content)
        issues.extend(validate_inventory_records(records))
        self.assertFalse([issue for issue in issues if issue['severity'] == 'error'])
        batch = create_import_batch(
            actor=self.admin, filename='incremental.xlsx', content=content,
            mode=ImportBatch.Mode.INCREMENTAL, records=records, issues=issues, summary=summary,
        )
        apply_inventory_import(actor=self.admin, batch=batch, records=records)
        self.stock.refresh_from_db()
        self.assertEqual(self.stock.quantity, 9)
        self.assertTrue(Asset.objects.filter(serial_number='IMPORT-SN-001').exists())
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.IMPORTED)

    def test_baseline_import_disables_missing_rows_only_in_listed_warehouse(self):
        missing_asset = Asset.objects.create(
            name='基准缺失资产', item_type=self.item_type, warehouse=self.warehouse,
            location=self.location, status=Asset.Status.IN_STOCK,
        )
        other_warehouse = Warehouse.objects.create(name='不在文件中的仓库', code='OTHER-WH')
        untouched_asset = Asset.objects.create(
            name='其他仓资产', item_type=self.item_type, warehouse=other_warehouse,
            status=Asset.Status.IN_STOCK,
        )
        included_asset = Asset.objects.create(
            name='基准保留资产', item_type=self.item_type, warehouse=self.warehouse,
            location=self.location, status=Asset.Status.IN_STOCK, serial_number='BASELINE-KEEP',
            asset_code='BASELINE-KEEP-ZC',
        )
        content = self._workbook([
            ['单件资产', included_asset.name, self.item_type.name, self.item_type.code, self.warehouse.name, self.location.code, included_asset.asset_code, included_asset.serial_number, '', 1, '件', 2, '在库', '基准保留'],
        ])
        records, issues, summary = parse_inventory_workbook(content)
        issues.extend(validate_inventory_records(records))
        batch = create_import_batch(
            actor=self.admin, filename='baseline.xlsx', content=content,
            mode=ImportBatch.Mode.BASELINE, records=records, issues=issues, summary=summary,
        )
        apply_inventory_import(actor=self.admin, batch=batch, records=records)
        missing_asset.refresh_from_db()
        untouched_asset.refresh_from_db()
        self.stock.refresh_from_db()
        self.assertEqual(missing_asset.status, Asset.Status.DISABLED)
        self.assertEqual(self.stock.quantity, 0)
        self.assertEqual(untouched_asset.status, Asset.Status.IN_STOCK)
        self.assertTrue(batch.issues.filter(code='baseline_asset_disabled').exists())

    def test_management_pages_expose_simplified_tools(self):
        self.client.force_login(self.admin)
        for path in ('/warehouse/items/new/', '/warehouse/low-stock/', '/warehouse/conversions/', '/warehouse/import/'):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, path)
        response = self.client.get('/warehouse/')
        self.assertContains(response, '管理工具')
        self.assertContains(response, '需要补货')


class InventoryCheckTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(username='check-admin', password='password123', email='')
        self.operator = User.objects.create_user(username='check-operator', password='password123', is_staff=True)
        self.operator.groups.add(Group.objects.get_or_create(name='warehouse_entry')[0])
        category = Category.objects.create(name='盘点测试分类', code='CHECK-CAT')
        self.item_type = ItemType.objects.create(name='盘点测试类型', code='CHECK-TYPE', category=category)
        self.warehouse = Warehouse.objects.create(name='盘点测试仓库', code='CHECK-WH')
        self.asset = Asset.objects.create(
            name='盘点测试资产', item_type=self.item_type, warehouse=self.warehouse,
            asset_code='CHECK-ZC-001', serial_number='CHECK-SN-001', qr_value='CHECK-QR-001',
        )
        self.stock = StockItem.objects.create(
            name='盘点测试耗材', code='CHECK-STOCK-001', item_type=self.item_type,
            warehouse=self.warehouse, quantity=5,
        )

    def test_check_freezes_snapshot_scans_and_reviews_difference(self):
        task = create_check_task(actor=self.admin, name='月度盘点', warehouse_id=self.warehouse.id)
        self.assertEqual(task.lines.count(), 2)
        start_check_task(task, self.operator)

        asset_scan = scan_check_value(task, self.operator, self.asset.qr_value)
        stock_scan = scan_check_value(task, self.operator, self.stock.code, quantity=4)
        self.assertEqual(asset_scan.result, 'resolved')
        self.assertEqual(stock_scan.result, 'resolved')

        submit_check_task(task, self.operator)
        task.refresh_from_db()
        self.assertEqual(task.status, InventoryCheckTask.Status.PENDING_REVIEW)
        stock_line = task.lines.get(stock_item=self.stock)
        self.assertEqual(stock_line.result, InventoryCheckLine.Result.SHORTAGE)
        self.assertEqual(Notification.objects.filter(recipient=self.admin, related_object_id=str(task.id)).count(), 1)

        _, adjustments = review_check_task(task, self.admin, '确认本次月度盘点差异')
        self.assertEqual(len(adjustments), 1)
        self.stock.refresh_from_db()
        task.refresh_from_db()
        self.assertEqual(self.stock.quantity, 4)
        self.assertEqual(task.status, InventoryCheckTask.Status.COMPLETED)
        self.assertTrue(InventoryCheckAdjustment.objects.filter(task=task, action='shortage').exists())
        self.assertTrue(OperationAuditLog.objects.filter(action='inventory_check_completed', target_id=task.id).exists())

    def test_unscanned_asset_becomes_shortage_and_lost_only_after_review(self):
        task = create_check_task(actor=self.admin, name='资产盘点', warehouse_id=self.warehouse.id)
        start_check_task(task, self.operator)
        submit_check_task(task, self.operator)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.IN_STOCK)

        review_check_task(task, self.admin, '确认资产未找到')
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.LOST)

    def test_status_change_after_snapshot_is_reviewed_as_status_difference(self):
        task = create_check_task(actor=self.admin, name='状态盘点', warehouse_id=self.warehouse.id)
        start_check_task(task, self.operator)
        self.asset.status = Asset.Status.REPAIRING
        self.asset.save(update_fields=['status', 'updated_at'])
        scan_check_value(task, self.operator, self.asset.asset_code)
        line = task.lines.get(asset=self.asset)
        self.assertEqual(line.result, InventoryCheckLine.Result.STATUS_MISMATCH)
        submit_check_task(task, self.operator)
        review_check_task(task, self.admin, '确认状态回到在库')
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.IN_STOCK)

    def test_only_inventory_editor_can_scan_and_only_admin_can_review(self):
        task = create_check_task(actor=self.admin, name='权限测试', warehouse_id=self.warehouse.id)
        response = self.client.post(f'/api/inventory/inventory-checks/{task.id}/start/')
        self.assertEqual(response.status_code, 401)

        self.client.force_login(self.operator)
        response = self.client.post(f'/api/inventory/inventory-checks/{task.id}/start/')
        self.assertEqual(response.status_code, 200, response.content)
        response = self.client.post(f'/api/inventory/inventory-checks/{task.id}/scan/', {
            'value': self.asset.asset_code,
        }, content_type='application/json')
        self.assertEqual(response.status_code, 200, response.content)

        submit_check_task(task, self.operator)
        response = self.client.post(f'/api/inventory/inventory-checks/{task.id}/review/', {}, content_type='application/json')
        self.assertEqual(response.status_code, 403)

        self.client.force_login(self.admin)
        response = self.client.get(f'/api/inventory/inventory-checks/{task.id}/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['line_total'], 2)

    def test_web_pages_expose_check_list_create_and_detail_flow(self):
        self.client.force_login(self.admin)
        response = self.client.get('/inventory-checks/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '新建盘点任务')

        response = self.client.post('/inventory-checks/new/', {
            'name': '网页盘点任务',
            'warehouse_id': str(self.warehouse.id),
            'note': '网页测试',
        })
        self.assertEqual(response.status_code, 302)
        task = InventoryCheckTask.objects.get(name='网页盘点任务')
        self.assertEqual(response.url, f'/inventory-checks/{task.id}/')

        self.client.force_login(self.operator)
        response = self.client.get(f'/inventory-checks/{task.id}/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '开始盘点')
        response = self.client.post(f'/inventory-checks/{task.id}/', {'action': 'start'})
        self.assertEqual(response.status_code, 302)
        response = self.client.post(f'/inventory-checks/{task.id}/', {
            'action': 'scan', 'scanned_value': self.asset.asset_code, 'quantity': '1',
        })
        self.assertEqual(response.status_code, 302)
        response = self.client.get(f'/inventory-checks/{task.id}/')
        self.assertContains(response, 'name="correction_note"')

    def test_admin_can_reopen_for_correction_and_export_all_check_sheets(self):
        task = create_check_task(actor=self.admin, name='导出与修正', warehouse_id=self.warehouse.id)
        start_check_task(task, self.operator)
        scan_check_value(task, self.operator, self.asset.asset_code)
        scan_check_value(task, self.operator, self.stock.code, quantity=4)
        submit_check_task(task, self.operator)

        self.client.force_login(self.admin)
        response = self.client.post(f'/inventory-checks/{task.id}/', {
            'action': 'reopen', 'reopen_note': '数量需要重新核对',
        })
        self.assertEqual(response.status_code, 302)
        task.refresh_from_db()
        self.assertEqual(task.status, InventoryCheckTask.Status.IN_PROGRESS)

        self.client.force_login(self.operator)
        stock_line = task.lines.get(stock_item=self.stock)
        response = self.client.post(f'/inventory-checks/{task.id}/', {
            'action': 'correct', 'line_id': str(stock_line.id), 'quantity': '5',
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '修改已录入的实盘数量必须填写修正原因')
        stock_line.refresh_from_db()
        self.assertEqual(stock_line.counted_quantity, 4)
        response = self.client.post(f'/inventory-checks/{task.id}/', {
            'action': 'correct', 'line_id': str(stock_line.id), 'quantity': '5', 'correction_note': '重新清点确认',
        })
        self.assertEqual(response.status_code, 302)
        stock_line.refresh_from_db()
        self.assertEqual(stock_line.result, InventoryCheckLine.Result.MATCHED)
        self.assertTrue(OperationAuditLog.objects.filter(action='inventory_check_line_corrected', target_id=stock_line.id).exists())
        submit_check_task(task, self.operator)

        self.client.force_login(self.admin)
        response = self.client.get(f'/inventory-checks/{task.id}/export/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        workbook = load_workbook(filename=__import__('io').BytesIO(response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ['任务摘要', '盘点明细', '扫描记录', '差异调整'])
        self.assertEqual(next(workbook['任务摘要'].iter_rows(min_row=2, max_row=2, values_only=True))[1], task.check_no)
        self.assertIn('任务编号', next(workbook['盘点明细'].iter_rows(min_row=1, max_row=1, values_only=True)))
        workbook.close()

        response = self.client.post(f'/inventory-checks/{task.id}/', {
            'action': 'review', 'review_note': '确认修正后的盘点结果',
        })
        self.assertEqual(response.status_code, 302)
        self.stock.refresh_from_db()
        self.assertEqual(self.stock.quantity, 5)


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class CompositeUnitAssemblyTests(TestCase):
    """成品设备：免审批的组装/换装/拆散 + 审批的出库/归还。"""

    def setUp(self):
        from apps.inventory.assembly_services import add_components, create_unit
        self._add_components = add_components
        self._create_unit = create_unit
        self.employee = User.objects.create_user(username='asm-employee', password='password123')
        self.operator = User.objects.create_user(username='asm-operator', password='password123')
        self.operator.groups.add(Group.objects.create(name='warehouse_staff'))
        category = Category.objects.create(name='组装测试分类', code='ASM-CAT')
        self.item_type = ItemType.objects.create(name='机器人整机', code='ASM-TYPE', category=category)
        self.warehouse = Warehouse.objects.create(name='组装测试仓', code='ASM-WH')
        self.assets = [
            Asset.objects.create(
                name=f'组件-{index}',
                item_type=self.item_type,
                warehouse=self.warehouse,
                status=Asset.Status.IN_STOCK,
            )
            for index in range(3)
        ]

    def _assembled_unit(self):
        unit = self._create_unit(actor=self.employee, name='宇树G1 演示整机', unit_code='DZSB-C-001')
        self._add_components(unit=unit, asset_ids=[a.id for a in self.assets], actor=self.employee)
        unit.refresh_from_db()
        return unit

    def test_create_shell_then_add_components(self):
        from apps.inventory.models import ComponentLink, CompositeUnit
        unit = self._create_unit(actor=self.employee, name='空壳成品', unit_code='DZSB-C-002')
        self.assertEqual(unit.status, CompositeUnit.Status.DRAFT)
        self.assertEqual(unit.qr_value, 'DZSB-C-002')

        self._add_components(unit=unit, asset_ids=[self.assets[0].id, self.assets[1].id], actor=self.employee)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.ASSEMBLED)
        self.assertEqual(ComponentLink.objects.filter(composite_unit=unit, disassembled_at__isnull=True).count(), 2)
        for asset in self.assets[:2]:
            asset.refresh_from_db()
            self.assertEqual(asset.status, Asset.Status.ASSEMBLED)
        self.assertTrue(OperationAuditLog.objects.filter(action='composite_components_added', target_id=unit.id).exists())

    def test_add_components_rejects_non_in_stock_asset(self):
        from apps.inventory.assembly_services import AssemblyError
        self.assets[0].status = Asset.Status.BORROWED
        self.assets[0].save(update_fields=['status', 'updated_at'])
        unit = self._create_unit(actor=self.employee, name='不应成功的组装')
        with self.assertRaises(AssemblyError) as ctx:
            self._add_components(unit=unit, asset_ids=[a.id for a in self.assets], actor=self.employee)
        self.assertIn('在库', str(ctx.exception))

    def test_component_cannot_be_reused_while_assembled(self):
        from apps.inventory.assembly_services import AssemblyError
        unit = self._assembled_unit()
        with self.assertRaises(AssemblyError) as ctx:
            self._add_components(unit=unit, asset_ids=[self.assets[0].id], actor=self.employee)
        self.assertIn('已组装', str(ctx.exception))

    def test_swap_component_on_unshipped_unit(self):
        from apps.inventory.assembly_services import swap_component
        from apps.inventory.models import ComponentLink
        unit = self._assembled_unit()
        spare = Asset.objects.create(name='备件', item_type=self.item_type, warehouse=self.warehouse, status=Asset.Status.IN_STOCK)
        swap_component(unit=unit, old_asset_id=self.assets[0].id, new_asset_id=spare.id, actor=self.employee, reason='电机异常')

        self.assets[0].refresh_from_db()
        spare.refresh_from_db()
        self.assertEqual(self.assets[0].status, Asset.Status.IN_STOCK)
        self.assertEqual(spare.status, Asset.Status.ASSEMBLED)
        active = ComponentLink.objects.filter(composite_unit=unit, disassembled_at__isnull=True)
        self.assertEqual(active.count(), 3)
        self.assertTrue(OperationAuditLog.objects.filter(action='composite_component_swapped', target_id=unit.id).exists())

    def test_swap_rejected_after_outbound(self):
        from apps.inventory.assembly_services import create_unit_outbound_request, swap_component, AssemblyError
        from apps.inventory.models import CompositeUnit
        from apps.workflow.services import approve_request
        unit = self._assembled_unit()
        request_obj = create_unit_outbound_request(
            unit=unit, request_type=WorkflowRequest.RequestType.BORROW, actor=self.employee,
            usage_location='客户现场',
        )
        approve_request(request_obj, self.operator)
        process_request(request_obj, self.operator)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.BORROWED)

        spare = Asset.objects.create(name='备件', item_type=self.item_type, warehouse=self.warehouse, status=Asset.Status.IN_STOCK)
        with self.assertRaises(AssemblyError) as ctx:
            swap_component(unit=unit, old_asset_id=self.assets[0].id, new_asset_id=spare.id, actor=self.employee)
        self.assertIn('冻结', str(ctx.exception))

    def test_outbound_and_return_sync_unit_status(self):
        from apps.inventory.assembly_services import create_unit_outbound_request, create_unit_return_request
        from apps.inventory.models import CompositeUnit
        from apps.workflow.services import approve_request
        unit = self._assembled_unit()
        request_obj = create_unit_outbound_request(
            unit=unit, request_type=WorkflowRequest.RequestType.BORROW, actor=self.employee,
            usage_location='客户现场',
        )
        approve_request(request_obj, self.operator)
        process_request(request_obj, self.operator)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.BORROWED)
        for asset in self.assets:
            asset.refresh_from_db()
            self.assertEqual(asset.status, Asset.Status.BORROWED)

        return_req = create_unit_return_request(unit=unit, actor=self.employee)
        approve_request(return_req, self.operator)
        process_request(return_req, self.operator)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.ASSEMBLED)
        for asset in self.assets:
            asset.refresh_from_db()
            self.assertEqual(asset.status, Asset.Status.ASSEMBLED)

    def test_disassembly_releases_components_and_voids_unit(self):
        from apps.inventory.assembly_services import disassemble_unit
        from apps.inventory.models import ComponentLink, CompositeUnit
        unit = self._assembled_unit()
        disassemble_unit(unit=unit, actor=self.employee, reason='项目结束')

        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.DISASSEMBLED)
        for asset in self.assets:
            asset.refresh_from_db()
            self.assertEqual(asset.status, Asset.Status.IN_STOCK)
        links = ComponentLink.objects.filter(composite_unit=unit)
        self.assertTrue(all(link.disassembled_at is not None for link in links))
        self.assertTrue(OperationAuditLog.objects.filter(action='composite_disassembled', target_id=unit.id).exists())

    def test_disassembled_asset_can_be_reassembled(self):
        from apps.inventory.assembly_services import disassemble_unit
        from apps.inventory.models import CompositeUnit
        unit = self._assembled_unit()
        disassemble_unit(unit=unit, actor=self.employee)

        unit2 = self._create_unit(actor=self.employee, name='第二次组装的成品')
        self._add_components(unit=unit2, asset_ids=[self.assets[0].id, self.assets[1].id], actor=self.employee)
        unit2.refresh_from_db()
        self.assertEqual(unit2.status, CompositeUnit.Status.ASSEMBLED)
        self.assets[0].refresh_from_db()
        self.assertEqual(self.assets[0].status, Asset.Status.ASSEMBLED)

    def test_unit_list_shows_active_component_count_not_total(self):
        """列表"组件数"必须只统计在装组件，与详情页口径一致（历史拆出不计入）。"""
        from apps.inventory.assembly_services import disassemble_unit
        unit = self._assembled_unit()  # 3 件在装
        disassemble_unit(unit=unit, actor=self.employee)  # 全部拆出
        # 拆散后成品默认从列表隐藏，需主动筛选已拆散状态
        unit.status = unit.__class__.Status.DRAFT
        unit.save(update_fields=['status', 'updated_at'])

        self.client.force_login(self.operator)
        response = self.client.get('/warehouse/composite-units/?status=draft')
        self.assertEqual(response.status_code, 200)
        units = list(response.context['units'])
        self.assertEqual(len(units), 1)
        # 在装数为 0，但 component_links 总数仍是 3
        self.assertEqual(units[0].active_component_count, 0)
        self.assertEqual(units[0].component_links.count(), 3)

    def test_unit_page_layout_filter_and_list_in_right_column(self):
        """筛选面板与成品列表必须同处右列容器，避免网格自动换行把列表排到左列。"""
        self.client.force_login(self.operator)
        response = self.client.get('/warehouse/composite-units/')
        self.assertEqual(response.status_code, 200)
        html = response.content.decode('utf-8')
        i_layout = html.find('unit-layout')
        i_sticky = html.find('unit-sticky')
        i_right = html.find('unit-right')
        i_filter = html.find('筛选成品设备')
        i_list = html.find('成品设备列表')
        # DOM 顺序：unit-layout > 左列(unit-sticky 新建) > 右列(unit-right) 内依次 筛选、列表
        self.assertTrue(i_layout < i_sticky < i_right < i_filter < i_list)

    def test_unit_borrow_via_request_create_page(self):
        """从新建申请页选成品借用：自动展开组件、关联 composite_unit、状态流转。"""
        unit = self._assembled_unit()
        self.client.force_login(self.employee)
        response = self.client.post('/requests/new/', {
            'request_type': WorkflowRequest.RequestType.BORROW,
            'reason': '现场演示',
            'recipient_name': '张三',
            'recipient_phone': '13800000000',
            'usage_location': '客户现场',
            'expected_return_date': (timezone.localdate() + __import__('datetime').timedelta(days=7)).isoformat(),
            'application_date': timezone.localdate().isoformat(),
            'unit_ids': [str(unit.id)],
            'action': 'submit',
        })
        self.assertEqual(response.status_code, 302)
        request_obj = WorkflowRequest.objects.filter(composite_unit=unit).latest('created_at')
        self.assertEqual(request_obj.request_type, WorkflowRequest.RequestType.BORROW)
        self.assertEqual(request_obj.status, WorkflowRequest.Status.PENDING)
        self.assertEqual(request_obj.lines.count(), 3)
        for link in unit.component_links.filter(disassembled_at__isnull=True):
            self.assertTrue(request_obj.lines.filter(asset_id=link.asset_id).exists())

    def test_unit_return_via_request_create_page(self):
        """借用出库后，从新建申请页发起归还：候选可找到、整台归还、状态回已组装。"""
        from apps.inventory.assembly_services import create_unit_outbound_request
        from apps.workflow.services import approve_request
        unit = self._assembled_unit()
        outbound = create_unit_outbound_request(
            unit=unit, request_type=WorkflowRequest.RequestType.BORROW, actor=self.employee,
            usage_location='客户现场',
        )
        approve_request(outbound, self.operator)
        process_request(outbound, self.operator)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.BORROWED)

        self.client.force_login(self.employee)
        response = self.client.post('/requests/new/', {
            'request_type': WorkflowRequest.RequestType.RETURN,
            'reason': '用完归还',
            'application_date': timezone.localdate().isoformat(),
            'unit_ids': [str(unit.id)],
            'action': 'submit',
        })
        self.assertEqual(response.status_code, 302)
        return_req = WorkflowRequest.objects.filter(composite_unit=unit, request_type=WorkflowRequest.RequestType.RETURN).latest('created_at')
        self.assertEqual(return_req.lines.count(), 3)
        approve_request(return_req, self.operator)
        process_request(return_req, self.operator)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.ASSEMBLED)

    def test_unit_search_shows_in_candidates(self):
        """申请页搜索能命中成品设备候选。"""
        unit = self._assembled_unit()
        self.client.force_login(self.employee)
        response = self.client.get('/requests/new/', {'q': '宇树G1'})
        self.assertContains(response, '成品设备')
        self.assertContains(response, unit.name)

    def test_return_variance_partial_receipt(self):
        """差异验收：归还时未勾选的组件保持借出，成品仍为借出，可再次归还剩余组件。"""
        from apps.inventory.assembly_services import create_unit_outbound_request, create_unit_return_request
        from apps.workflow.services import approve_request, process_request
        unit = self._assembled_unit()
        outbound = create_unit_outbound_request(
            unit=unit, request_type=WorkflowRequest.RequestType.BORROW, actor=self.employee,
            usage_location='客户现场',
        )
        approve_request(outbound, self.operator)
        process_request(outbound, self.operator)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.BORROWED)

        return_req = create_unit_return_request(unit=unit, actor=self.employee)
        approve_request(return_req, self.operator)
        lines = list(return_req.lines.order_by('created_at'))
        received_ids = [str(lines[0].id), str(lines[1].id)]  # 少收 1 件
        process_request(return_req, self.operator, return_received_line_ids=received_ids)

        # 收回的 2 件回到已组装，未收回的 1 件仍借出，成品仍为已借出
        lines[0].asset.refresh_from_db(); lines[1].asset.refresh_from_db(); lines[2].asset.refresh_from_db()
        self.assertEqual(lines[0].asset.status, Asset.Status.ASSEMBLED)
        self.assertEqual(lines[1].asset.status, Asset.Status.ASSEMBLED)
        self.assertEqual(lines[2].asset.status, Asset.Status.BORROWED)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.BORROWED)
        self.assertTrue(OperationAuditLog.objects.filter(action='return_variance', target_id=return_req.id).exists())

        # 剩余 1 件可再次发起归还并全部收回，成品回到已组装
        return_req2 = create_unit_return_request(unit=unit, actor=self.employee)
        self.assertEqual(return_req2.lines.count(), 1)
        approve_request(return_req2, self.operator)
        process_request(return_req2, self.operator)
        unit.refresh_from_db()
        self.assertEqual(unit.status, CompositeUnit.Status.ASSEMBLED)
        lines[2].asset.refresh_from_db()
        self.assertEqual(lines[2].asset.status, Asset.Status.ASSEMBLED)

    def test_delete_empty_draft_unit(self):
        """待组装的空壳成品可直接删除，并写审计。"""
        from apps.inventory.assembly_services import delete_unit
        from apps.inventory.models import CompositeUnit
        unit = self._create_unit(actor=self.employee, name='误建的测试成品', unit_code='DZSB-C-999')
        label = delete_unit(unit=unit, actor=self.employee, reason='测试')
        self.assertEqual(label, 'DZSB-C-999')
        self.assertFalse(CompositeUnit.objects.filter(pk=unit.pk).exists())
        self.assertTrue(OperationAuditLog.objects.filter(action='composite_unit_deleted').exists())

    def test_delete_rejects_assembled_unit(self):
        """有组件的成品（已组装）禁止删除。"""
        from apps.inventory.assembly_services import AssemblyError, delete_unit
        unit = self._assembled_unit()
        with self.assertRaises(AssemblyError) as ctx:
            delete_unit(unit=unit, actor=self.employee)
        self.assertIn('已组装', str(ctx.exception))

    def test_delete_rejects_unit_with_history(self):
        """拆散后有历史组件履历的成品禁止删除。"""
        from apps.inventory.assembly_services import AssemblyError, delete_unit, disassemble_unit
        unit = self._assembled_unit()
        disassemble_unit(unit=unit, actor=self.employee)
        with self.assertRaises(AssemblyError) as ctx:
            delete_unit(unit=unit, actor=self.employee)
        self.assertIn('履历', str(ctx.exception))

    def test_force_delete_rejects_non_admin(self):
        """非管理员不能强制删除。"""
        from apps.inventory.assembly_services import AssemblyError, delete_unit, disassemble_unit
        unit = self._assembled_unit()
        disassemble_unit(unit=unit, actor=self.employee)
        with self.assertRaises(AssemblyError) as ctx:
            delete_unit(unit=unit, actor=self.employee, force=True)
        self.assertIn('管理员', str(ctx.exception))
        self.assertTrue(CompositeUnit.objects.filter(pk=unit.pk).exists())

    def test_force_delete_unit_with_history_as_admin(self):
        """管理员可强删有历史履历的已拆散成品。"""
        from apps.inventory.assembly_services import delete_unit, disassemble_unit
        admin = User.objects.create_user(username='unit-admin', password='password123', is_superuser=True)
        unit = self._assembled_unit()
        disassemble_unit(unit=unit, actor=self.employee)
        label = delete_unit(unit=unit, actor=admin, reason='清理测试数据', force=True)
        self.assertEqual(label, str(unit))
        self.assertFalse(CompositeUnit.objects.filter(pk=unit.pk).exists())
        self.assertFalse(ComponentLink.objects.filter(composite_unit_id=unit.pk).exists())
        audit = OperationAuditLog.objects.filter(action='composite_unit_deleted').latest('created_at')
        self.assertTrue(audit.before_data['force'])

    def test_force_delete_assembled_unit_releases_components(self):
        """管理员强删已组装成品时，组件自动释放回在库。"""
        from apps.inventory.assembly_services import delete_unit
        admin = User.objects.create_user(username='unit-admin2', password='password123', is_superuser=True)
        unit = self._assembled_unit()
        asset_ids = [link.asset_id for link in unit.component_links.all()]
        delete_unit(unit=unit, actor=admin, reason='测试清理', force=True)
        self.assertFalse(CompositeUnit.objects.filter(pk=unit.pk).exists())
        for asset_id in asset_ids:
            asset = Asset.objects.get(pk=asset_id)
            self.assertEqual(asset.status, Asset.Status.IN_STOCK)

    def _bom_workbook(self, rows):
        from io import BytesIO
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(['成品名称', '成品编码', '型号', '组件资产编码', '装入备注'])
        for row in rows:
            ws.append(row)
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def test_bom_import_creates_units_and_components(self):
        """BOM Excel 批量导入：同一成品多行合并、组件装入、状态流转。"""
        from apps.inventory.assembly_import import import_unit_bom
        from apps.inventory.models import CompositeUnit
        codes = []
        for index, asset in enumerate(self.assets):
            asset.asset_code = f'BOM-{index:04d}'
            asset.save(update_fields=['asset_code'])
            codes.append(asset.asset_code)
        content = self._bom_workbook([
            ['G1 整机', 'G1-001', 'G1-Pro', codes[0], '主控'],
            ['G1 整机', 'G1-001', 'G1-Pro', codes[1], ''],
            ['G2 整机', 'G2-001', '', codes[2], '雷达'],
        ])
        units = import_unit_bom(actor=self.employee, content=content)
        self.assertEqual(len(units), 2)
        g1 = CompositeUnit.objects.get(name='G1 整机')
        self.assertEqual(g1.unit_code, 'G1-001')
        self.assertEqual(g1.status, CompositeUnit.Status.ASSEMBLED)
        self.assertEqual(g1.component_links.filter(disassembled_at__isnull=True).count(), 2)
        for asset in self.assets[:2]:
            asset.refresh_from_db()
            self.assertEqual(asset.status, Asset.Status.ASSEMBLED)

    def test_bom_import_rolls_back_on_error(self):
        """任一行失败时整体回滚，不产生半成品数据。"""
        from apps.inventory.assembly_import import BomImportError, import_unit_bom
        from apps.inventory.models import CompositeUnit
        self.assets[0].asset_code = 'BOM-OK'
        self.assets[0].save(update_fields=['asset_code'])
        content = self._bom_workbook([
            ['G1 整机', 'G1-001', '', 'BOM-OK', ''],
            ['G1 整机', 'G1-001', '', 'NOT-EXIST', ''],
        ])
        with self.assertRaises(BomImportError) as ctx:
            import_unit_bom(actor=self.employee, content=content)
        self.assertIn('NOT-EXIST', str(ctx.exception))
        self.assertFalse(CompositeUnit.objects.filter(name='G1 整机').exists())
        self.assets[0].refresh_from_db()
        self.assertEqual(self.assets[0].status, Asset.Status.IN_STOCK)
