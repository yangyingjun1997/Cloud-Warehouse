from __future__ import annotations

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import InventoryCheckAdjustment, InventoryCheckLine, InventoryCheckTask


HEADER_FILL = PatternFill('solid', fgColor='1F4E78')
HEADER_FONT = Font(color='FFFFFF', bold=True)


def _excel_datetime(value):
    if value and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value or ''


def _display(obj, method_name):
    return getattr(obj, method_name)() if obj else ''


def _format_sheet(sheet, widths):
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)


def export_check_task_xlsx(task: InventoryCheckTask) -> BytesIO:
    task = InventoryCheckTask.objects.select_related(
        'warehouse', 'location', 'created_by', 'started_by', 'reviewed_by', 'reopened_by',
    ).prefetch_related(
        'lines__item_type', 'lines__asset', 'lines__stock_item',
        'lines__expected_warehouse', 'lines__expected_location',
        'lines__observed_warehouse', 'lines__observed_location', 'lines__counted_by',
        'scans__asset', 'scans__stock_item', 'scans__operator',
        'scans__observed_warehouse', 'scans__observed_location',
        'adjustments__asset', 'adjustments__stock_item', 'adjustments__reviewed_by',
        'adjustments__before_warehouse', 'adjustments__after_warehouse',
        'adjustments__before_location', 'adjustments__after_location',
    ).get(pk=task.pk)

    workbook = Workbook()
    summary = workbook.active
    summary.title = '任务摘要'
    summary.append(['字段', '内容'])
    summary_rows = [
        ('任务编号', task.check_no),
        ('任务名称', task.name),
        ('任务状态', task.get_status_display()),
        ('盘点仓库', task.warehouse.name),
        ('盘点库位', task.location.code if task.location else '全仓库'),
        ('创建人', task.created_by.username if task.created_by else ''),
        ('创建时间', _excel_datetime(task.created_at)),
        ('开始人', task.started_by.username if task.started_by else ''),
        ('开始时间', _excel_datetime(task.started_at)),
        ('提交时间', _excel_datetime(task.submitted_at)),
        ('最近退回人', task.reopened_by.username if task.reopened_by else ''),
        ('最近退回时间', _excel_datetime(task.reopened_at)),
        ('退回原因', task.reopen_note),
        ('复核人', task.reviewed_by.username if task.reviewed_by else ''),
        ('完成时间', _excel_datetime(task.completed_at)),
        ('任务说明', task.note),
        ('明细总数', task.lines.count()),
        ('差异项目数', task.lines.exclude(result=InventoryCheckLine.Result.MATCHED).count()),
        ('调整台账数', task.adjustments.count()),
    ]
    for row in summary_rows:
        summary.append(list(row))
    _format_sheet(summary, [24, 60])

    details = workbook.create_sheet('盘点明细')
    details.append([
        '任务编号', '物品类型', '对象类型', '仓库系统资产编号', '公司 ERP 资产编码', '资产名称', '生产厂家 SN', '物料编码',
        '账面数量', '实盘数量', '差异数量', '账面状态', '盘点结果', '账面仓库', '账面库位',
        '现场仓库', '现场库位', '盘点人', '盘点时间', '明细备注',
    ])
    for line in task.lines.all():
        is_asset = bool(line.asset_id)
        expected = line.expected_quantity
        counted = line.counted_quantity
        details.append([
            task.check_no,
            line.item_type.name if line.item_type else '未归类',
            '单件资产' if is_asset else '耗材配件',
            line.asset.system_asset_no if line.asset else line.expected_asset_code,
            line.asset.asset_code if line.asset else '',
            line.asset.name if line.asset else '',
            line.asset.serial_number if line.asset else line.expected_serial_number,
            line.stock_item.code if line.stock_item else '',
            expected,
            counted if counted is not None else '',
            counted - expected if counted is not None else '',
            _display(line.asset, 'get_status_display') if line.asset else line.expected_status,
            line.get_result_display(),
            line.expected_warehouse.name if line.expected_warehouse else '',
            line.expected_location.code if line.expected_location else '',
            line.observed_warehouse.name if line.observed_warehouse else '',
            line.observed_location.code if line.observed_location else '',
            line.counted_by.username if line.counted_by else '',
            _excel_datetime(line.counted_at),
            line.note,
        ])
    _format_sheet(details, [18, 18, 14, 22, 24, 24, 22, 22, 12, 12, 12, 16, 16, 20, 16, 20, 16, 16, 20, 30])

    scans = workbook.create_sheet('扫描记录')
    scans.append(['时间', '扫描内容', '扫描结果', '仓库系统资产编号', '公司 ERP 资产编码', '耗材配件', '数量', '操作人', '现场仓库', '现场库位', '备注'])
    for scan in task.scans.all():
        scans.append([
            _excel_datetime(scan.created_at), scan.scanned_value, scan.get_result_display(),
            scan.asset.system_asset_no if scan.asset else '', scan.asset.asset_code if scan.asset else '', scan.stock_item.name if scan.stock_item else '',
            scan.quantity, scan.operator.username if scan.operator else '',
            scan.observed_warehouse.name if scan.observed_warehouse else '',
            scan.observed_location.code if scan.observed_location else '', scan.note,
        ])
    _format_sheet(scans, [20, 30, 16, 22, 24, 24, 10, 16, 20, 16, 30])

    adjustments = workbook.create_sheet('差异调整')
    adjustments.append([
        '执行时间', '调整类型', '仓库系统资产编号', '公司 ERP 资产编码', '耗材配件', '调整前数量', '调整后数量', '调整前状态',
        '调整后状态', '调整前仓库', '调整后仓库', '调整前库位', '调整后库位', '复核人', '复核说明',
    ])
    for adjustment in task.adjustments.all():
        adjustments.append([
            _excel_datetime(adjustment.executed_at), adjustment.get_action_display(),
            adjustment.asset.system_asset_no if adjustment.asset else '',
            adjustment.asset.asset_code if adjustment.asset else '',
            adjustment.stock_item.name if adjustment.stock_item else '',
            adjustment.before_quantity if adjustment.before_quantity is not None else '',
            adjustment.after_quantity if adjustment.after_quantity is not None else '',
            adjustment.before_status, adjustment.after_status,
            adjustment.before_warehouse.name if adjustment.before_warehouse else '',
            adjustment.after_warehouse.name if adjustment.after_warehouse else '',
            adjustment.before_location.code if adjustment.before_location else '',
            adjustment.after_location.code if adjustment.after_location else '',
            adjustment.reviewed_by.username if adjustment.reviewed_by else '', adjustment.review_note,
        ])
    _format_sheet(adjustments, [20, 18, 22, 24, 24, 14, 14, 16, 16, 20, 20, 16, 16, 16, 36])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
