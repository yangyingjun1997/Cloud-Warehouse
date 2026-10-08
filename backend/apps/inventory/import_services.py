from __future__ import annotations

import hashlib
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from apps.common.audit import record_operation
from apps.common.utils import build_code

from .models import Asset, Category, ImportBatch, ImportIssue, ItemType, Location, StockItem, Supplier, Warehouse


HEADER_ALIASES = {
    'management': {'管理方式', '属于资产还是耗材', '资产/耗材', '资产或耗材'},
    'name': {'物品名称', '资产名称', '产品名称', '名称'},
    'item_type': {'物品类型', '类型'},
    'item_type_code': {'物品类型编码', '类型编码'},
    'warehouse': {'仓库', '所属仓库'},
    'location': {'库位', '所在库位'},
    'system_asset_no': {'仓库系统资产编号', '本系统资产编号', '系统资产编号'},
    # “系统编号” is retained here for older files where it meant the ERP
    # code. New templates use the explicit column names above.
    'asset_code': {'公司ERP资产编码', '公司 ERP 资产编码', '行政资产编码', '我司资产编码', '资产编号', '系统编号'},
    'serial_number': {'生产厂家SN/序列号', '生产厂家 SN / 序列号', '生产厂家SN', '厂家出厂SN序列号', '厂家SN', '原厂序列号', '序列号', 'SN'},
    'manufacturer_barcode': {'生产厂家条码/二维码内容', '生产厂家条码/二维码', '厂家条码/二维码', '厂家条码', '厂家二维码'},
    'qr_value': {'本系统二维码内容', '系统二维码内容', '我司二维码', '系统二维码'},
    'stock_code': {'耗材配件编码', '物料编码', '库存编码'},
    'quantity': {'库存数量', '数量', '当前库存'},
    'unit': {'单位', '计量单位'},
    'safety_stock': {'最低库存提醒值', '安全库存'},
    'manufacturer': {'生产厂家', '厂家', '制造商'},
    'model': {'型号', '规格型号'},
    'supplier': {'供应商'},
    'unit_cost': {'单价', '单价（元）', '采购单价'},
    'status': {'状态', '资产状态'},
    'remarks': {'备注', '说明'},
}

MANAGEMENT_VALUES = {
    '单件资产': 'asset', '资产': 'asset', '单件': 'asset', '逐件管理': 'asset',
    '耗材配件': 'stock', '耗材': 'stock', '配件': 'stock', '数量物料': 'stock', '按数量管理': 'stock',
}

STATUS_VALUES = {
    **{label: value for value, label in Asset.Status.choices},
    **{value: value for value, _ in Asset.Status.choices},
}


class InventoryImportError(ValueError):
    pass


def _text(value) -> str:
    return '' if value is None else str(value).strip()


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _number(value, *, integer=False, default=None):
    if value in (None, ''):
        return default
    try:
        number = Decimal(str(value).replace(',', '').strip())
        if number < 0:
            raise ValueError
        if integer and number != number.to_integral_value():
            raise ValueError
        return int(number) if integer else number
    except (InvalidOperation, ValueError):
        raise InventoryImportError(f'“{value}”不是有效的非负数。')


def _header_map(values):
    normalized = {_text(value).rstrip('*').strip(): index for index, value in enumerate(values) if _text(value)}
    return {
        key: next((normalized[alias] for alias in aliases if alias in normalized), None)
        for key, aliases in HEADER_ALIASES.items()
    }


def _cell(row, headers, key):
    index = headers.get(key)
    return row[index] if index is not None and index < len(row) else None


def _issue(sheet, row, severity, code, message, source):
    return {
        'sheet_name': sheet,
        'row_number': row,
        'severity': severity,
        'code': code,
        'message': message[:255],
        'source_data': {key: _json_value(value) for key, value in source.items()},
    }


def parse_inventory_workbook(content: bytes):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise InventoryImportError('无法读取该 Excel，请确认文件为有效的 .xlsx 文件。') from exc

    records, issues = [], []
    source_identifiers = Counter()
    for worksheet in workbook.worksheets:
        if worksheet.title.strip() in {'填写说明', '导入说明'}:
            continue
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration:
            continue
        headers = _header_map(header_values)
        missing = [key for key in ('management', 'name', 'item_type', 'warehouse') if headers.get(key) is None]
        if missing:
            issues.append(_issue(
                worksheet.title, 1, 'error', 'missing_columns',
                '缺少必需列：管理方式、物品名称、物品类型、仓库。', {},
            ))
            continue

        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, '') for value in row):
                continue
            source = {
                key: _cell(row, headers, key)
                for key in HEADER_ALIASES
                if headers.get(key) is not None
            }
            management_text = _text(source.get('management'))
            management = MANAGEMENT_VALUES.get(management_text)
            name = _text(source.get('name'))
            item_type_name = _text(source.get('item_type'))
            warehouse_name = _text(source.get('warehouse'))
            if not management:
                issues.append(_issue(worksheet.title, row_number, 'error', 'unknown_management', f'无法识别管理方式“{management_text}”。', source))
                continue
            if not name or not item_type_name or not warehouse_name:
                issues.append(_issue(worksheet.title, row_number, 'error', 'missing_required_value', '物品名称、物品类型和仓库不能为空。', source))
                continue
            try:
                quantity = _number(source.get('quantity'), integer=True, default=1 if management == 'asset' else None)
                safety_stock = _number(source.get('safety_stock'), integer=True, default=0)
                unit_cost = _number(source.get('unit_cost'), default=None)
            except InventoryImportError as exc:
                issues.append(_issue(worksheet.title, row_number, 'error', 'invalid_number', str(exc), source))
                continue
            status_text = _text(source.get('status'))
            if management == 'asset' and status_text and status_text not in STATUS_VALUES:
                issues.append(_issue(worksheet.title, row_number, 'error', 'unknown_status', f'无法识别资产状态“{status_text}”。', source))
                continue
            if management == 'stock' and quantity is None:
                issues.append(_issue(worksheet.title, row_number, 'error', 'missing_quantity', '耗材配件必须填写库存数量。', source))
                continue
            if management == 'asset' and quantity != 1:
                issues.append(_issue(worksheet.title, row_number, 'error', 'asset_quantity', '单件资产每行数量必须为 1；多件请拆成多行。', source))
                continue

            record = {
                'sheet_name': worksheet.title,
                'row_number': row_number,
                'source': source,
                'management': management,
                'name': name[:120],
                'item_type': item_type_name[:120],
                'item_type_code': _text(source.get('item_type_code'))[:50],
                'warehouse': warehouse_name[:120],
                'location': _text(source.get('location'))[:120],
                'system_asset_no': _text(source.get('system_asset_no'))[:32],
                'asset_code': _text(source.get('asset_code'))[:80],
                'serial_number': _text(source.get('serial_number'))[:120],
                'manufacturer_barcode': _text(source.get('manufacturer_barcode'))[:255],
                'qr_value': _text(source.get('qr_value'))[:255],
                'stock_code': _text(source.get('stock_code'))[:50],
                'quantity': quantity,
                'unit': _text(source.get('unit'))[:20] or '件',
                'safety_stock': safety_stock,
                'manufacturer': _text(source.get('manufacturer'))[:120],
                'model': _text(source.get('model'))[:120],
                'supplier': _text(source.get('supplier'))[:120],
                'unit_cost': unit_cost,
                'status': STATUS_VALUES.get(status_text, Asset.Status.IN_STOCK),
                'remarks': _text(source.get('remarks')),
            }
            identifiers = []
            if management == 'asset':
                identifiers = [
                    f'system_asset_no:{record["system_asset_no"]}',
                    f'asset_code:{record["asset_code"]}',
                    f'sn:{record["serial_number"]}',
                    f'barcode:{record["manufacturer_barcode"]}',
                    f'qr:{record["qr_value"]}',
                ]
            elif record['stock_code']:
                identifiers = [f'stock_code:{record["stock_code"]}']
            for identifier in (value for value in identifiers if not value.endswith(':')):
                source_identifiers[identifier] += 1
                record.setdefault('identifiers', []).append(identifier)
            records.append(record)
    workbook.close()

    duplicate_identifiers = {key for key, count in source_identifiers.items() if count > 1}
    safe_records = []
    for record in records:
        duplicates = duplicate_identifiers.intersection(record.get('identifiers', []))
        if duplicates:
            issues.append(_issue(
                record['sheet_name'], record['row_number'], 'error', 'duplicate_identifier',
                f'文件内唯一标识重复：{", ".join(sorted(duplicates))}。', record['source'],
            ))
            continue
        safe_records.append(record)

    summary = {
        'row_count': len(records),
        'valid_row_count': len(safe_records),
        'asset_count': sum(row['management'] == 'asset' for row in safe_records),
        'stock_item_count': sum(row['management'] == 'stock' for row in safe_records),
        'error_count': sum(issue['severity'] == 'error' for issue in issues),
        'warning_count': sum(issue['severity'] == 'warning' for issue in issues),
        'sheets': workbook.sheetnames,
    }
    return safe_records, issues, summary


def validate_inventory_records(records):
    issues = []
    for record in records:
        try:
            warehouse = _warehouse(record)
            _location(record, warehouse)
            by_code = ItemType.objects.filter(code=record['item_type_code']).first() if record['item_type_code'] else None
            by_names = ItemType.objects.filter(name=record['item_type'])
            if by_names.count() > 1 and not by_code:
                raise InventoryImportError('存在多个同名物品类型，请填写物品类型编码。')
            by_name = by_names.first()
            if by_code and by_name and by_code.pk != by_name.pk:
                raise InventoryImportError('物品类型名称和类型编码匹配到不同记录。')
            if record['management'] == 'asset':
                _find_asset(record)
            elif record['stock_code']:
                stock_by_code = StockItem.objects.filter(code=record['stock_code']).first()
                if stock_by_code and by_code and stock_by_code.item_type_id != by_code.id:
                    raise InventoryImportError('耗材配件编码与物品类型编码指向不同记录。')
            elif by_code or by_name:
                matches = StockItem.objects.filter(
                    item_type=by_code or by_name,
                    warehouse=warehouse,
                    location=_location(record, warehouse),
                )
                if matches.count() > 1:
                    raise InventoryImportError('匹配到多个耗材配件，请填写耗材配件编码。')
        except InventoryImportError as exc:
            issues.append(_issue(
                record['sheet_name'], record['row_number'], 'error',
                'database_conflict', str(exc), record['source'],
            ))
    return issues


def create_import_batch(*, actor, filename, content, mode, records, issues, summary):
    batch = ImportBatch.objects.create(
        source_filename=filename[:255],
        source_sha256=hashlib.sha256(content).hexdigest(),
        mode=mode,
        status=ImportBatch.Status.PREVIEWED,
        summary=summary,
        created_by=actor,
    )
    ImportIssue.objects.bulk_create([ImportIssue(batch=batch, **issue) for issue in issues])
    return batch


def _find_asset(record):
    matches = set()
    for field in ('system_asset_no', 'asset_code', 'serial_number', 'manufacturer_barcode', 'qr_value'):
        value = record.get(field)
        if value:
            matches.update(Asset.objects.filter(**{field: value}).values_list('id', flat=True))
    if len(matches) > 1:
        raise InventoryImportError('仓库系统编号、公司 ERP 编码、生产厂家 SN、条码或二维码分别匹配到不同资产。')
    return Asset.objects.filter(pk=next(iter(matches))).first() if matches else None


def _warehouse(record):
    warehouse = Warehouse.objects.filter(name=record['warehouse']).first() or Warehouse.objects.filter(code=record['warehouse']).first()
    if not warehouse:
        raise InventoryImportError(f'仓库“{record["warehouse"]}”不存在，请先在基础资料中创建。')
    return warehouse


def _location(record, warehouse):
    if not record['location']:
        return None
    location = Location.objects.filter(warehouse=warehouse, name=record['location']).first() or Location.objects.filter(
        warehouse=warehouse, code=record['location']
    ).first()
    if not location:
        raise InventoryImportError(f'仓库“{warehouse.name}”中不存在库位“{record["location"]}”。')
    return location


def _item_type(record, category):
    by_code = ItemType.objects.filter(code=record['item_type_code']).first() if record['item_type_code'] else None
    by_name = ItemType.objects.filter(name=record['item_type']).first()
    if by_code and by_name and by_code.pk != by_name.pk:
        raise InventoryImportError('物品类型名称和类型编码匹配到不同记录。')
    item_type = by_code or by_name
    if item_type:
        if record['management'] == 'asset' and record['safety_stock']:
            item_type.safety_stock = record['safety_stock']
            item_type.save(update_fields=['safety_stock', 'updated_at'])
        return item_type
    return ItemType.objects.create(
        name=record['item_type'],
        code=record['item_type_code'] or build_code('TYPE')[:50],
        category=category,
        is_serialized=record['management'] == 'asset',
        unit=record['unit'],
        safety_stock=record['safety_stock'],
    )


def _supplier(name, actor):
    if not name:
        return None
    supplier = Supplier.objects.filter(name=name).first()
    if supplier:
        return supplier
    supplier = Supplier.objects.create(code=build_code('SUPPLIER')[:80], name=name)
    record_operation(
        actor,
        'supplier_created_from_inventory_import',
        supplier,
        summary=f'Excel 导入自动建立供应商 {name}',
        after={'name': name, 'code': supplier.code},
    )
    return supplier


@transaction.atomic
def apply_inventory_import(*, actor, batch, records):
    if batch.status != ImportBatch.Status.PREVIEWED:
        raise InventoryImportError('该导入批次已经处理，不能重复执行。')
    if batch.issues.filter(severity=ImportIssue.Severity.ERROR).exists():
        raise InventoryImportError('预检存在无法导入的问题，请修正 Excel 后重新上传。')

    category, _ = Category.objects.get_or_create(code='DEFAULT', defaults={'name': '默认分类'})
    imported_assets = imported_stock = 0
    created_count = updated_count = disabled_count = 0
    included_asset_ids, included_stock_ids, represented_warehouse_ids = set(), set(), set()

    for record in records:
        warehouse = _warehouse(record)
        location = _location(record, warehouse)
        item_type = _item_type(record, category)
        supplier_ref = _supplier(record['supplier'], actor)
        represented_warehouse_ids.add(warehouse.id)
        if record['management'] == 'asset':
            asset = _find_asset(record)
            is_create = asset is None
            if is_create:
                asset = Asset(
                    system_asset_no=record['system_asset_no'] or None,
                    asset_code=record['asset_code'] or None,
                )
            before = {
                'name': asset.name if not is_create else '',
                'warehouse_id': asset.warehouse_id if not is_create else None,
                'location_id': asset.location_id if not is_create else None,
                'status': asset.status if not is_create else '',
            }
            asset.name = record['name']
            asset.item_type = item_type
            asset.warehouse = warehouse
            asset.location = location
            asset.status = record['status']
            asset.source_import_batch = batch
            asset.supplier_ref = supplier_ref
            for field in ('serial_number', 'manufacturer_barcode', 'qr_value', 'manufacturer', 'model', 'supplier', 'remarks'):
                if record[field] or is_create:
                    setattr(asset, field, record[field])
            if record['unit_cost'] is not None:
                asset.purchase_amount = record['unit_cost']
            asset.save()
            included_asset_ids.add(asset.id)
            imported_assets += 1
            created_count += int(is_create)
            updated_count += int(not is_create)
            record_operation(actor, 'inventory_import_asset_created' if is_create else 'inventory_import_asset_updated', asset, summary=f'{batch.get_mode_display()}：{asset.system_asset_no}', before=before, after={'name': asset.name, 'warehouse_id': warehouse.id, 'location_id': location.id if location else None, 'status': asset.status})
        else:
            stock_item = StockItem.objects.filter(code=record['stock_code']).first() if record['stock_code'] else None
            if stock_item is None:
                matches = StockItem.objects.filter(item_type=item_type, warehouse=warehouse, location=location)
                if matches.count() > 1:
                    raise InventoryImportError(f'第 {record["row_number"]} 行匹配到多个耗材配件，请填写耗材配件编码。')
                stock_item = matches.first()
            is_create = stock_item is None
            if is_create:
                stock_item = StockItem(code=record['stock_code'] or build_code('STOCK'))
            before_quantity = stock_item.quantity if not is_create else 0
            stock_item.name = record['name']
            stock_item.item_type = item_type
            stock_item.warehouse = warehouse
            stock_item.location = location
            stock_item.quantity = record['quantity']
            stock_item.unit = record['unit']
            stock_item.safety_stock = record['safety_stock']
            stock_item.supplier = record['supplier']
            stock_item.supplier_ref = supplier_ref
            stock_item.unit_cost = record['unit_cost']
            stock_item.source_import_batch = batch
            stock_item.save()
            included_stock_ids.add(stock_item.id)
            imported_stock += 1
            created_count += int(is_create)
            updated_count += int(not is_create)
            record_operation(actor, 'inventory_import_stock_created' if is_create else 'inventory_import_stock_updated', stock_item, summary=f'{batch.get_mode_display()}：{stock_item.code}', before={'quantity': before_quantity}, after={'quantity': stock_item.quantity})

    if batch.mode == ImportBatch.Mode.BASELINE and represented_warehouse_ids:
        missing_assets = Asset.objects.select_for_update().filter(
            warehouse_id__in=represented_warehouse_ids,
            status=Asset.Status.IN_STOCK,
        ).exclude(pk__in=included_asset_ids)
        for asset in missing_assets:
            before_status = asset.status
            asset.status = Asset.Status.DISABLED
            asset.save(update_fields=['status', 'updated_at'])
            ImportIssue.objects.create(batch=batch, sheet_name='全量基准', row_number=0, severity=ImportIssue.Severity.INFO, code='baseline_asset_disabled', message=f'{asset.asset_code or f"未编号·{asset.name}"} 未出现在基准文件中，已标记为停用。', source_data={'asset_id': str(asset.id)}, resolved=True, resolution_note='由全量基准导入自动处理。')
            record_operation(actor, 'inventory_baseline_asset_disabled', asset, summary=f'{asset.asset_code or f"未编号·{asset.name}"} 未出现在全量基准中', before={'status': before_status}, after={'status': asset.status, 'import_batch_id': batch.id})
            disabled_count += 1
        missing_stock = StockItem.objects.select_for_update().filter(
            warehouse_id__in=represented_warehouse_ids, quantity__gt=0,
        ).exclude(pk__in=included_stock_ids)
        for stock_item in missing_stock:
            before_quantity = stock_item.quantity
            stock_item.quantity = 0
            stock_item.save(update_fields=['quantity', 'updated_at'])
            ImportIssue.objects.create(batch=batch, sheet_name='全量基准', row_number=0, severity=ImportIssue.Severity.INFO, code='baseline_stock_zeroed', message=f'{stock_item.code} 未出现在基准文件中，库存已归零。', source_data={'stock_item_id': str(stock_item.id)}, resolved=True, resolution_note='由全量基准导入自动处理。')
            record_operation(actor, 'inventory_baseline_stock_zeroed', stock_item, summary=f'{stock_item.code} 未出现在全量基准中', before={'quantity': before_quantity}, after={'quantity': 0, 'import_batch_id': batch.id})
            disabled_count += 1

    batch.status = ImportBatch.Status.IMPORTED
    batch.imported_asset_count = imported_assets
    batch.imported_stock_item_count = imported_stock
    batch.summary = {
        **batch.summary,
        'created_count': created_count,
        'updated_count': updated_count,
        'baseline_disabled_or_zeroed_count': disabled_count,
    }
    batch.save(update_fields=['status', 'imported_asset_count', 'imported_stock_item_count', 'summary', 'updated_at'])
    record_operation(actor, 'inventory_import_applied', batch, summary=f'{batch.get_mode_display()}完成：处理 {imported_assets + imported_stock} 行', after=batch.summary)
    return batch


def build_inventory_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '库存导入'
    headers = [
        '管理方式*', '物品名称*', '物品类型*', '物品类型编码', '仓库*', '库位',
        '仓库系统资产编号', '公司 ERP 资产编码', '生产厂家 SN / 序列号', '生产厂家条码 / 二维码内容', '本系统二维码内容',
        '耗材配件编码', '库存数量', '单位', '最低库存提醒值', '厂家', '型号',
        '供应商', '单价（元）', '状态', '备注',
    ]
    sheet.append(headers)
    example_rows = [
        ['单件资产', '四足机器人本体', '机器狗', 'ROBOT-DOG', '市场库存', 'A-01', '', '', 'DOG-SN-2026-001', 'DOG-MFG-001', '', '', 1, '台', 2, '示例机器人厂家', 'D1-Pro', '示例供应商甲', 68000, '在库', '示例数据，正式填写时请替换或删除'],
        ['单件资产', '三维激光雷达', '雷达', 'LIDAR', '研发库存', 'R-02', '', '', 'LIDAR-SN-2026-008', 'LIDAR-MFG-008', '', '', 1, '台', 3, '示例雷达厂家', 'L32', '示例供应商乙', 12500, '在库', '单件资产每台占一行'],
        ['单件资产', '边缘计算工控机', '工控机', 'IPC', '研发库存', 'N-01', 'AST-000003', 'DZSB-D7-21-0001-600001', 'IPC-SN-0001', '', '', '', 1, '台', 1, '示例计算机厂家', 'i7-32G', '示例供应商丙', 7600, '在库', '仓库系统编号由系统生成；公司 ERP 编码可留空'],
        ['耗材配件', '六类成品网线', '网络配件', 'NET-PART', '研发库存', 'B-01', '', '', '', '', '', 'CABLE-CAT6-003M', 30, '根', 8, '示例线缆厂家', 'CAT6-3M', '示例供应商乙', 18.5, '', '耗材配件按总数量管理'],
        ['耗材配件', '航空插头组件', '连接器', 'CONNECTOR', '市场库存', 'C-03', '', '', '', '', '', 'CONNECTOR-GX16-4', 50, '套', 10, '示例连接器厂家', 'GX16-4', '示例供应商甲', 12.8, '', '耗材配件按总数量管理'],
    ]
    for row in example_rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176E68')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:U{sheet.max_row}'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(_text(cell.value)) for cell in column) + 3, 28)
    notes = workbook.create_sheet('填写说明')
    notes.append(['项目', '说明'])
    notes.append(['必填列', '管理方式、物品名称、物品类型、仓库；模板主表中列名带 *。'])
    notes.append(['条件必填', '耗材配件必须填写库存数量；单件资产数量可留空，系统按 1 件处理。'])
    notes.append(['仓库系统资产编号', '系统自动生成 AST-000001 这类不可修改的内部编号；已有资产更新时可填写它进行精确匹配，新资产留空即可。'])
    notes.append(['公司 ERP 资产编码', '公司行政/ERP 分配的资产编码，有就填写，没有就留空；它可以后续补录。'])
    notes.append(['生产厂家 SN', '生产厂家提供的 SN/序列号；有就填写，没有就留空，不要用内部编号冒充。'])
    notes.append(['系统二维码', '本系统二维码默认只编码仓库系统资产编号；本列原则上留空，系统会自动生成。生产厂家条码/二维码请填写到对应的生产厂家列。这些都是公司内部字段。'])
    notes.append(['再次导入', '首次导入后建议保留并导出仓库系统资产编号；系统编号、ERP 编码、生产厂家 SN、生产厂家条码和本系统二维码均可用于匹配。'])
    notes.append(['可选列', '库位、物品类型编码、厂家条码/二维码、厂家、型号、供应商、单价、状态、备注均可留空。'])
    notes.append(['可选列默认值', '库位为空表示未定位；物品类型编码为空时同名类型唯一即可；单位默认为“件”；最低库存默认为 0；资产状态默认为“在库”。'])
    notes.append(['编码生成', '耗材配件编码留空且无法匹配已有记录时自动生成；新物品类型编码留空时自动生成；供应商填写名称但未建档时会自动建立供应商编码。'])
    notes.append(['管理方式', '只能填写“单件资产”或“耗材配件”。'])
    notes.append(['增量导入', '新增不存在的数据，并更新唯一匹配的数据；文件中没有的记录保持不变。'])
    notes.append(['全量基准导入', '仅处理文件中出现的仓库；缺失的在库资产会停用，缺失的耗材配件库存会归零。'])
    notes.append(['唯一匹配', '单件资产依次使用仓库系统资产编号、公司 ERP 编码、生产厂家 SN、生产厂家条码和本系统二维码；耗材配件优先使用耗材配件编码。'])
    notes.append(['示例数据', '“库存导入”中的 5 行仅用于演示填写方式，正式导入前请替换或删除。'])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
