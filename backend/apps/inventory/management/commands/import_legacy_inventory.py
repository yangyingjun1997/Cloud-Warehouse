"""Preview and import the approved legacy inventory workbook.

The default mode intentionally creates only a review batch and issue list.  Use
--apply only after the preview has been checked in the management interface.
"""

from __future__ import annotations

import hashlib
import ipaddress
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.inventory.models import (
    Asset,
    AssetNetworkEndpoint,
    Category,
    ImportBatch,
    ImportIssue,
    ItemType,
    StockItem,
    Warehouse,
)


WAREHOUSE_CONFIG = {
    '市场库存': 'MARKET',
    '研发库存': 'RND',
}

STATUS_MAP = {
    '市场库存': {
        '在库空闲': Asset.Status.IN_STOCK,
        '出库借用': Asset.Status.BORROWED,
        '内部借用库': Asset.Status.BORROWED,
        '出库出售': Asset.Status.SOLD,
        '返厂维修': Asset.Status.REPAIRING,
        '销售预定': Asset.Status.PENDING_OUT,
        '退货': Asset.Status.RETURNED_TO_VENDOR,
        '报废': Asset.Status.SCRAPPED,
        '待质检': Asset.Status.PENDING_INSPECTION,
    },
    '研发库存': {
        '未出库': Asset.Status.IN_STOCK,
        '已出库': Asset.Status.ISSUED,
        '返厂': Asset.Status.REPAIRING,
        '归还': Asset.Status.IN_STOCK,
        '报废': Asset.Status.SCRAPPED,
    },
}

PLACEHOLDER_IDENTIFIERS = {'', '/', '-', '无', '未知', 'n/a', 'na', '本批没有序列号'}


class Command(BaseCommand):
    help = '预检或导入旧版市场/研发库存 Excel（默认仅生成复核批次）'

    def add_arguments(self, parser):
        parser.add_argument('workbook', type=Path, help='Excel 文件路径')
        parser.add_argument('--apply', action='store_true', help='确认将可导入资产写入正式库')
        parser.add_argument('--user', help='执行导入的系统账号（用于追溯）')

    def handle(self, *args, **options):
        path = options['workbook'].expanduser().resolve()
        if not path.is_file():
            raise CommandError(f'找不到 Excel 文件: {path}')
        if path.suffix.lower() not in {'.xlsx', '.xlsm'}:
            raise CommandError('仅支持 .xlsx 或 .xlsm 文件。')

        user = self._get_user(options.get('user'))
        asset_records, stock_records, issues, summary = self._parse_workbook(path)
        batch = ImportBatch.objects.create(
            source_filename=path.name,
            source_sha256=self._sha256(path),
            status=ImportBatch.Status.PREVIEWED,
            summary=summary,
            created_by=user,
        )
        self._save_issues(batch, issues)

        if options['apply']:
            imported_assets, imported_stock_items, imported_endpoints, apply_issues = self._apply_records(
                batch,
                asset_records,
                stock_records,
            )
            self._save_issues(batch, apply_issues)
            batch.status = ImportBatch.Status.IMPORTED
            batch.imported_asset_count = imported_assets
            batch.imported_stock_item_count = imported_stock_items
            batch.imported_network_endpoint_count = imported_endpoints
            batch.summary = {**summary, 'apply_issue_count': len(apply_issues)}
            batch.save(update_fields=[
                'status',
                'imported_asset_count',
                'imported_stock_item_count',
                'imported_network_endpoint_count',
                'summary',
                'updated_at',
            ])
            self.stdout.write(self.style.SUCCESS(
                f'导入完成：资产 {imported_assets} 条，库存物料 {imported_stock_items} 条，'
                f'网络接口 {imported_endpoints} 条，请在后台查看复核问题。'
            ))
        else:
            self.stdout.write(self.style.WARNING(
                f'预检完成：单件资产 {len(asset_records)} 条，库存物料 {len(stock_records)} 条，'
                f'需复核 {len(issues)} 条。'
                f' 批次 ID: {batch.id}。确认后使用 --apply 重新导入。'
            ))

    def _get_user(self, username):
        if not username:
            return None
        try:
            return get_user_model().objects.get(username=username)
        except get_user_model().DoesNotExist as exc:
            raise CommandError(f'账号不存在: {username}') from exc

    def _parse_workbook(self, path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        asset_records, stock_records, issues = [], [], []
        for sheet_name in WAREHOUSE_CONFIG:
            if sheet_name not in workbook.sheetnames:
                issues.append(self._issue(sheet_name, 0, 'error', 'missing_sheet', f'缺少工作表“{sheet_name}”。', {}))
                continue
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            try:
                headers = self._headers(next(rows))
            except StopIteration:
                issues.append(self._issue(sheet_name, 0, 'error', 'empty_sheet', '工作表为空。', {}))
                continue

            product_name = ''
            source_serials = set()
            source_codes = set()
            for row_number, row in enumerate(rows, start=2):
                source = self._source_row(headers, row)
                product_name = self._text(source.get('产品名称')) or product_name
                if not any(source.values()):
                    continue
                if not product_name:
                    issues.append(self._issue(sheet_name, row_number, 'error', 'missing_product', '缺少产品名称，无法归类。', source))
                    continue

                serial = self._text(source.get('序列号'))
                internal_code = self._text(source.get('资产编号') or source.get('八维通SN'))
                status_source = self._text(source.get('状态'))
                ip_value = serial if self._is_ip(serial) else ''
                usable_serial = '' if self._is_placeholder(serial) or ip_value else serial
                usable_code = '' if self._is_placeholder(internal_code) else internal_code

                if not (usable_serial or usable_code or ip_value):
                    if self._is_summary_row(sheet_name, source):
                        stock_records.append({
                            'sheet_name': sheet_name,
                            'row_number': row_number,
                            'source': source,
                            'name': product_name,
                            'quantity': self._summary_quantity(sheet_name, source),
                        })
                        continue
                    issues.append(self._issue(
                        sheet_name,
                        row_number,
                        'warning',
                        'unidentified_row',
                        '无有效序列号、旧资产编号或 IP，已转入复核清单。',
                        source,
                    ))
                    continue
                if self._is_serial_range(usable_serial):
                    issues.append(self._issue(sheet_name, row_number, 'warning', 'serial_range', '序列号为范围值，无法对应单件资产，已转入复核清单。', source))
                    continue
                if usable_serial and usable_serial in source_serials:
                    issues.append(self._issue(sheet_name, row_number, 'error', 'duplicate_serial', '同一工作簿序列号重复，未纳入自动导入。', source))
                    continue
                if usable_code and usable_code in source_codes:
                    issues.append(self._issue(sheet_name, row_number, 'error', 'duplicate_internal_code', '同一工作簿旧资产编号重复，未纳入自动导入。', source))
                    continue
                source_serials.add(usable_serial)
                source_codes.add(usable_code)

                mapped_status = STATUS_MAP[sheet_name].get(status_source, Asset.Status.IN_STOCK)
                if status_source and status_source not in STATUS_MAP[sheet_name]:
                    issues.append(self._issue(sheet_name, row_number, 'warning', 'unknown_status', f'未识别状态“{status_source}”，暂按“在库”导入。', source))
                if ip_value:
                    issues.append(self._issue(sheet_name, row_number, 'info', 'ip_in_serial_column', 'IP 已转存为网络接口记录，厂家 SN 保留为空。', source))
                asset_records.append({
                    'sheet_name': sheet_name,
                    'row_number': row_number,
                    'source': source,
                    'name': product_name,
                    'serial_number': usable_serial,
                    'asset_code': usable_code,
                    'ip_address': ip_value,
                    'status': mapped_status,
                })

        ip_counts = Counter(record['ip_address'] for record in asset_records if record['ip_address'])
        safe_asset_records = []
        for record in asset_records:
            if record['ip_address'] and ip_counts[record['ip_address']] > 1:
                issues.append(self._issue(
                    record['sheet_name'],
                    record['row_number'],
                    'error',
                    'duplicate_ip',
                    '工作簿内存在重复 IP，无法确定对应单件，未纳入自动导入。',
                    record['source'],
                ))
                continue
            safe_asset_records.append(record)

        asset_product_keys = {(record['sheet_name'], record['name']) for record in safe_asset_records}
        standalone_stock_records = []
        for record in stock_records:
            if (record['sheet_name'], record['name']) in asset_product_keys:
                issues.append({
                    **self._issue(
                        record['sheet_name'],
                        record['row_number'],
                        'info',
                        'summary_represented_by_assets',
                        '产品汇总数已由同仓库单件资产反映，不重复创建数量库存。',
                        record['source'],
                    ),
                    'resolved': True,
                    'resolution_note': '已按单件资产导入，汇总行仅作核对。',
                })
                continue
            standalone_stock_records.append(record)

        summary = {
            'candidate_asset_count': len(safe_asset_records),
            'candidate_stock_item_count': len(standalone_stock_records),
            'issue_count': len(issues),
            'status_counts': dict(Counter(record['status'] for record in safe_asset_records)),
            'sheets': list(WAREHOUSE_CONFIG),
        }
        return safe_asset_records, standalone_stock_records, issues, summary

    def _apply_records(self, batch, asset_records, stock_records):
        issues = []
        imported_assets = imported_stock_items = imported_endpoints = 0
        with transaction.atomic():
            warehouses = {
                name: Warehouse.objects.get_or_create(name=name, defaults={'code': code})[0]
                for name, code in WAREHOUSE_CONFIG.items()
            }
            category, _ = Category.objects.get_or_create(name='历史导入', defaults={'code': 'LEGACY'})
            for record in stock_records:
                stock_code = self._stock_item_code(record['sheet_name'], record['name'])
                if StockItem.objects.filter(code=stock_code).exists():
                    issues.append(self._issue(record['sheet_name'], record['row_number'], 'warning', 'existing_stock_item', '系统内已存在同仓库的同名库存物料，未覆盖原数量。', record['source']))
                    continue
                StockItem.objects.create(
                    name=record['name'][:120],
                    code=stock_code,
                    unit='个',
                    quantity=record['quantity'],
                    item_type=self._item_type(category, record['name'], is_serialized=False),
                    warehouse=warehouses[record['sheet_name']],
                    source_import_batch=batch,
                )
                imported_stock_items += 1
            for record in asset_records:
                source = record['source']
                serial = record['serial_number']
                requested_code = record['asset_code']
                if serial and Asset.objects.filter(serial_number=serial).exists():
                    issues.append(self._issue(record['sheet_name'], record['row_number'], 'error', 'existing_serial', '系统内已存在相同厂家 SN，已跳过。', source))
                    continue
                if requested_code and Asset.objects.filter(asset_code=requested_code).exists():
                    issues.append(self._issue(record['sheet_name'], record['row_number'], 'error', 'existing_asset_code', '系统内已存在相同资产编号，已跳过。', source))
                    continue

                item_type = self._item_type(category, record['name'])
                asset = Asset.objects.create(
                    asset_code=requested_code,
                    name=record['name'],
                    item_type=item_type,
                    manufacturer=self._text(source.get('厂家')),
                    model=self._text(source.get('型号')),
                    serial_number=serial,
                    received_date=self._date(source.get('到货日期')),
                    purchase_order_no=self._reference(source),
                    warehouse=warehouses[record['sheet_name']],
                    status=record['status'],
                    source_import_batch=batch,
                    remarks=self._remarks(record['sheet_name'], source),
                )
                imported_assets += 1
                if record['ip_address']:
                    AssetNetworkEndpoint.objects.create(asset=asset, ip_address=record['ip_address'], is_primary=True)
                    imported_endpoints += 1
        return imported_assets, imported_stock_items, imported_endpoints, issues

    def _item_type(self, category, name, *, is_serialized=True):
        code = f'LEGACY-{hashlib.sha1(name.encode("utf-8")).hexdigest()[:12].upper()}'
        return ItemType.objects.get_or_create(
            category=category,
            name=name[:120],
            defaults={'code': code, 'is_serialized': is_serialized, 'unit': '个'},
        )[0]

    @staticmethod
    def _stock_item_code(sheet_name, name):
        value = f'{sheet_name}:{name}'.encode('utf-8')
        return f'STOCK-{hashlib.sha1(value).hexdigest()[:12].upper()}'

    @staticmethod
    def _is_summary_row(sheet_name, source):
        quantity_field = '库存剩余' if sheet_name == '市场库存' else '库存数量'
        return source.get('产品名称') is not None and source.get(quantity_field) is not None and not source.get('状态')

    def _summary_quantity(self, sheet_name, source):
        quantity_field = '库存剩余' if sheet_name == '市场库存' else '库存数量'
        value = source.get(quantity_field)
        if isinstance(value, int):
            return max(value, 0)
        matched = re.search(r'\d+', self._text(value))
        return int(matched.group()) if matched else 0

    @staticmethod
    def _headers(row):
        return {str(value).strip(): index for index, value in enumerate(row) if value is not None and str(value).strip()}

    def _source_row(self, headers, row):
        return {
            header: self._json_value(row[index]) if index < len(row) else ''
            for header, index in headers.items()
        }

    @staticmethod
    def _text(value):
        if value is None:
            return ''
        return str(value).strip()

    def _is_placeholder(self, value):
        return self._text(value).lower() in PLACEHOLDER_IDENTIFIERS

    def _is_ip(self, value):
        try:
            ipaddress.ip_address(self._text(value))
            return True
        except ValueError:
            return False

    @staticmethod
    def _is_serial_range(value):
        return bool(re.fullmatch(r'\d+\s*-\s*\d+', value))

    @staticmethod
    def _date(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return date.fromisoformat(str(value))
        except (TypeError, ValueError):
            return None

    def _reference(self, source):
        return self._text(source.get('OA采购订单审批号') or source.get('ERP同步情况'))[:120]

    def _remarks(self, sheet_name, source):
        fields = ['客户名称', '销卖出库日期', '物流单号', '出货日期', '出货类型', '领用人', '归还日期', '用途', '备注']
        values = [f'{field}: {self._text(source.get(field))}' for field in fields if self._text(source.get(field))]
        return f'历史导入/{sheet_name}; ' + '; '.join(values)

    @staticmethod
    def _json_value(value):
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return value

    @staticmethod
    def _issue(sheet_name, row_number, severity, code, message, source_data):
        return {
            'sheet_name': sheet_name,
            'row_number': row_number,
            'severity': severity,
            'code': code,
            'message': message,
            'source_data': source_data,
        }

    @staticmethod
    def _save_issues(batch, issues):
        ImportIssue.objects.bulk_create([ImportIssue(batch=batch, **issue) for issue in issues])

    @staticmethod
    def _sha256(path):
        digest = hashlib.sha256()
        with path.open('rb') as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b''):
                digest.update(chunk)
        return digest.hexdigest()
