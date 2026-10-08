from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.apps import apps
from rest_framework.test import APIClient
from django.test import TestCase
from unittest.mock import patch
from datetime import date, timedelta
from uuid import uuid4
from importlib import import_module

from django.db import OperationalError
from django.utils import timezone

from apps.accounts.models import Department
from apps.inventory.models import Asset, Category, Customer, ItemType, Location, Project, StockItem, StockItemHolding, Warehouse, WorkOrder
from apps.notifications.models import Notification
from apps.common.models import OperationAuditLog

from .models import (
    ApprovalLog,
    ApprovalTask,
    InventoryReservation,
    InventoryTransaction,
    LoanExtensionRequest,
    OfflineOperation,
    PurchaseOrder,
    PurchaseReceipt,
    PurchaseReceiptLine,
    PurchaseRequest,
    StockItemLoan,
    WorkflowRequest,
)
from apps.common.database import retry_database_operation
from .services import (
    WorkflowError, close_request_and_release_reservations,
    create_loan_extension_request, expire_pending_reservations,
    quick_process_request, review_loan_extension,
    approve_purchase_request, create_purchase_order, place_purchase_order,
    receive_purchase_order,
)
from .offline_services import apply_offline_operation, receive_offline_operation, reject_offline_operation

User = get_user_model()


class OfflineOperationTests(TestCase):
    def setUp(self):
        self.operator = User.objects.create_user(username='offline-operator', password='password123')
        self.operator.groups.add(Group.objects.create(name='warehouse_staff'))
        self.warehouse = Warehouse.objects.create(name='离线测试仓', code='OFFLINE-WH')
        category = Category.objects.create(name='离线测试分类', code='OFFLINE-CATEGORY')
        item_type = ItemType.objects.create(name='离线测试资产', code='OFFLINE-ASSET', category=category)
        self.asset = Asset.objects.create(name='离线测试机器人', item_type=item_type, warehouse=self.warehouse)
        self.payload = {
            'request_type': ['borrow'],
            'asset_ids': [str(self.asset.id)],
            'recipient_name': ['测试接收人'],
            'recipient_phone': ['13800000000'],
            'usage_location': ['客户现场'],
            'reason': ['断网期间紧急出库'],
            'contact_source': ['电话'],
            'transaction_date': [timezone.localdate().isoformat()],
            'expected_return_date': [(timezone.localdate() + timedelta(days=7)).isoformat()],
        }

    def test_receive_is_idempotent_and_audited(self):
        client_id = uuid4()
        first, created = receive_offline_operation(
            operator=self.operator,
            client_operation_id=str(client_id),
            client_created_at=timezone.now().isoformat(),
            payload=self.payload,
        )
        second, created_again = receive_offline_operation(
            operator=self.operator,
            client_operation_id=str(client_id),
            client_created_at=timezone.now().isoformat(),
            payload=self.payload,
        )
        self.assertTrue(created)
        self.assertFalse(created_again)
        self.assertEqual(first.id, second.id)
        self.assertEqual(OfflineOperation.objects.count(), 1)
        self.assertTrue(OperationAuditLog.objects.filter(action='offline_operation_received', target_id=first.id).exists())

    def test_sync_endpoint_returns_server_review_status(self):
        self.client.force_login(self.operator)
        client_id = str(uuid4())
        form_data = {
            key: values if len(values) > 1 else values[0]
            for key, values in self.payload.items()
        }
        response = self.client.post('/offline/operations/sync/', {
            **form_data,
            'client_operation_id': client_id,
            'client_created_at': timezone.now().isoformat(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['status'], OfflineOperation.Status.PENDING)
        self.assertEqual(OfflineOperation.objects.get().payload['asset_ids'], [str(self.asset.id)])

    def test_review_applies_exactly_one_inventory_transaction(self):
        operation, _ = receive_offline_operation(
            operator=self.operator,
            client_operation_id=str(uuid4()),
            client_created_at=timezone.now().isoformat(),
            payload=self.payload,
        )
        applied = apply_offline_operation(operation, self.operator)
        applied_again = apply_offline_operation(operation, self.operator)

        self.asset.refresh_from_db()
        self.assertEqual(applied.status, OfflineOperation.Status.APPLIED)
        self.assertEqual(applied.id, applied_again.id)
        self.assertEqual(self.asset.status, Asset.Status.BORROWED)
        self.assertEqual(InventoryTransaction.objects.filter(request=applied.resulting_request).count(), 1)

    def test_changed_asset_state_becomes_conflict_without_inventory_write(self):
        operation, _ = receive_offline_operation(
            operator=self.operator,
            client_operation_id=str(uuid4()),
            client_created_at=timezone.now().isoformat(),
            payload=self.payload,
        )
        self.asset.status = Asset.Status.SOLD
        self.asset.save(update_fields=['status', 'updated_at'])

        result = apply_offline_operation(operation, self.operator)

        self.assertEqual(result.status, OfflineOperation.Status.CONFLICT)
        self.assertIn('无法出库', result.error_message)
        self.assertEqual(InventoryTransaction.objects.count(), 0)
        self.assertEqual(WorkflowRequest.objects.count(), 0)

    def test_rejected_operation_cannot_be_applied(self):
        operation, _ = receive_offline_operation(
            operator=self.operator,
            client_operation_id=str(uuid4()),
            client_created_at=timezone.now().isoformat(),
            payload=self.payload,
        )
        reject_offline_operation(operation, self.operator, '离线登记信息不完整')
        with self.assertRaisesMessage(WorkflowError, '已被拒绝'):
            apply_offline_operation(operation, self.operator)

    def test_reference_selection_backfills_legacy_text_fields(self):
        customer = Customer.objects.create(code='CUSTOMER-1', name='示例客户')
        project = Project.objects.create(code='PROJECT-1', name='示例项目', customer=customer)
        work_order = WorkOrder.objects.create(code='WO-1', name='现场维修工单', project=project)
        request_obj = WorkflowRequest.objects.create(
            request_type=WorkflowRequest.RequestType.ISSUE,
            applicant=self.operator,
            customer_ref=customer,
            project_ref=project,
            work_order_ref=work_order,
        )
        self.assertEqual(request_obj.recipient_company, customer.name)
        self.assertEqual(request_obj.project_code, project.code)
        self.assertEqual(request_obj.work_order_no, work_order.code)


class WorkflowApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.department = Department.objects.create(name='售后部', code='AFTERSALES')
        self.employee = User.objects.create_user(username='employee', password='password123')
        self.operator = User.objects.create_user(username='operator', password='password123', is_staff=True)
        self.operator.groups.add(Group.objects.create(name='warehouse_staff'))
        self.warehouse = Warehouse.objects.create(name='售后仓', code='AFTER-WH')
        self.destination_warehouse = Warehouse.objects.create(name='备用仓', code='SPARE-WH')
        self.destination_location = Location.objects.create(
            warehouse=self.destination_warehouse,
            name='A 区 01 架',
            code='A-01',
        )
        category = Category.objects.create(name='机器人载荷', code='PAYLOAD')
        item_type = ItemType.objects.create(name='测试载荷', code='PAYLOAD-TEST', category=category)
        self.asset = Asset.objects.create(
            name='测试载荷 01',
            item_type=item_type,
            warehouse=self.warehouse,
            department=self.department,
        )
        self.stock_item = StockItem.objects.create(
            name='测试耗材',
            code='CONSUMABLE-TEST',
            item_type=item_type,
            warehouse=self.warehouse,
            quantity=5,
        )

    def _create_request(self, request_type='borrow'):
        return self._create_request_with_lines(
            request_type,
            [{'asset': str(self.asset.id), 'quantity': 1}],
        )

    def _create_request_with_lines(self, request_type, lines, **extra):
        traceability = {}
        if request_type in {
            WorkflowRequest.RequestType.BORROW,
            WorkflowRequest.RequestType.ISSUE,
            WorkflowRequest.RequestType.SALE,
        }:
            traceability = {
                'usage_location': '售后测试区',
            }
        self.client.force_authenticate(self.employee)
        response = self.client.post('/api/workflow/requests/', {
            'request_type': request_type,
            'reason': '现场售后测试',
            'lines': lines,
            **traceability,
            **extra,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        return response.data['id']

    def _approve_and_process(self, request_id):
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        self.client.force_authenticate(self.operator)
        self.client.post(f'/api/workflow/requests/{request_id}/approve/', {}, format='json')
        response = self.client.post(f'/api/workflow/requests/{request_id}/process/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        return response

    def test_sqlite_lock_is_retried_before_workflow_operation_fails(self):
        attempts = []

        def operation():
            attempts.append(True)
            if len(attempts) < 3:
                raise OperationalError('database is locked')
            return 'done'

        with patch('apps.common.database.time.sleep'):
            self.assertEqual(retry_database_operation(operation), 'done')
        self.assertEqual(len(attempts), 3)

    def test_borrow_submit_approve_process_updates_asset_and_audit(self):
        request_id = self._create_request()

        response = self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['status'], WorkflowRequest.Status.PENDING)
        self.assertEqual(response.data['lines'][0]['asset_code'], self.asset.asset_code)
        self.assertEqual(response.data['lines'][0]['asset_name'], self.asset.name)
        self.assertEqual(ApprovalTask.objects.count(), 1)
        self.assertEqual(Notification.objects.filter(recipient=self.operator).count(), 1)

        response = self.client.post(f'/api/workflow/requests/{request_id}/approve/', {}, format='json')
        self.assertEqual(response.status_code, 403)

        self.client.force_authenticate(self.operator)
        response = self.client.post(f'/api/workflow/requests/{request_id}/approve/', {'comment': '同意'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['status'], WorkflowRequest.Status.WAITING_WAREHOUSE)

        response = self.client.post(f'/api/workflow/requests/{request_id}/process/', {'comment': '已扫码出库'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['status'], WorkflowRequest.Status.DONE)
        self.assertRegex(response.data['inventory_no'], r'^IO-\d{8}-\d{6}$')

        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.BORROWED)
        self.assertEqual(self.asset.location_state, Asset.LocationState.OUT_ON_LOAN)
        self.assertEqual(self.asset.availability_state, Asset.AvailabilityState.UNAVAILABLE)
        self.assertEqual(self.asset.current_holder_id, self.employee.id)
        request = WorkflowRequest.objects.get(pk=request_id)
        transaction = InventoryTransaction.objects.get(request=request)
        self.assertRegex(request.request_no, r'^REQ-\d{8}-\d{6}$')
        self.assertRegex(request.inventory_no, r'^IO-\d{8}-\d{6}$')
        self.assertEqual(transaction.business_date, request.transaction_date)
        self.assertEqual(ApprovalLog.objects.filter(request_id=request_id).count(), 3)

    def test_new_request_numbers_increment_by_application_date(self):
        first = WorkflowRequest.objects.create(
            applicant=self.employee,
            request_type=WorkflowRequest.RequestType.BORROW,
            application_date=date(2026, 8, 27),
        )
        second = WorkflowRequest.objects.create(
            applicant=self.employee,
            request_type=WorkflowRequest.RequestType.BORROW,
            application_date=date(2026, 8, 27),
        )

        self.assertEqual(first.request_no, 'REQ-20260827-000001')
        self.assertEqual(second.request_no, 'REQ-20260827-000002')

    def test_superuser_can_approve_and_process_workflow_requests(self):
        request_id = self._create_request()
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        administrator = User.objects.create_superuser(username='workflow-admin', password='password123', email='')
        self.client.force_authenticate(administrator)

        response = self.client.post(f'/api/workflow/requests/{request_id}/approve/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        response = self.client.post(f'/api/workflow/requests/{request_id}/process/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)

    def test_return_flow_requires_current_holder_and_puts_asset_back_in_stock(self):
        self.asset.current_holder = self.employee
        self.asset.status = Asset.Status.BORROWED
        self.asset.save(update_fields=['current_holder', 'status', 'updated_at'])
        request_id = self._create_request('return')

        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        self.client.force_authenticate(self.operator)
        self.client.post(f'/api/workflow/requests/{request_id}/approve/', {}, format='json')
        response = self.client.post(f'/api/workflow/requests/{request_id}/process/', {}, format='json')

        self.assertEqual(response.status_code, 200, response.data)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.IN_STOCK)
        self.assertIsNone(self.asset.current_holder_id)

    def test_stock_item_borrow_and_return_updates_inventory_and_personal_holding(self):
        borrow_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.BORROW,
            [{'stock_item': str(self.stock_item.id), 'quantity': 2}],
        )
        self._approve_and_process(borrow_id)
        self.stock_item.refresh_from_db()
        self.assertEqual(self.stock_item.quantity, 3)
        self.assertEqual(
            StockItemHolding.objects.get(stock_item=self.stock_item, holder=self.employee).quantity,
            2,
        )

        return_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.RETURN,
            [{'stock_item': str(self.stock_item.id), 'quantity': 2}],
        )
        self._approve_and_process(return_id)
        self.stock_item.refresh_from_db()
        self.assertEqual(self.stock_item.quantity, 5)
        self.assertEqual(
            StockItemHolding.objects.get(stock_item=self.stock_item, holder=self.employee).quantity,
            0,
        )
        transactions = InventoryTransaction.objects.filter(stock_item=self.stock_item).order_by('created_at')
        self.assertEqual(transactions.count(), 2)
        self.assertEqual(transactions.first().stock_quantity_before, 5)
        self.assertEqual(transactions.first().stock_quantity_after, 3)

    def test_employee_cannot_return_more_stock_than_personal_borrowing_balance(self):
        StockItemHolding.objects.create(
            stock_item=self.stock_item,
            holder=self.employee,
            quantity=1,
        )
        request_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.RETURN,
            [{'stock_item': str(self.stock_item.id), 'quantity': 2}],
        )

        response = self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')

        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('本人可归还数量不足', response.data['detail'])

    def test_transfer_moves_asset_to_target_warehouse_and_records_transaction(self):
        request_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.TRANSFER,
            [{'asset': str(self.asset.id), 'quantity': 1}],
            target_warehouse=str(self.destination_warehouse.id),
            target_location=str(self.destination_location.id),
        )
        self._approve_and_process(request_id)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.warehouse_id, self.destination_warehouse.id)
        self.assertEqual(self.asset.location_id, self.destination_location.id)
        transaction = InventoryTransaction.objects.get(asset=self.asset)
        self.assertEqual(transaction.source_warehouse_id, self.warehouse.id)
        self.assertEqual(transaction.target_warehouse_id, self.destination_warehouse.id)

    def test_external_recipient_sale_records_snapshot_and_final_asset_status(self):
        request_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.SALE,
            [{'asset': str(self.asset.id), 'quantity': 1}],
            recipient_name='客户张工',
            recipient_company='客户公司',
            recipient_phone='13800138000',
        )
        self._approve_and_process(request_id)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.SOLD)
        self.assertEqual(self.asset.external_holder_name, '客户张工')
        transaction = InventoryTransaction.objects.get(asset=self.asset)
        self.assertEqual(transaction.holder_company, '客户公司')

    def test_repair_and_scrap_change_asset_status_and_create_transactions(self):
        repair_id = self._create_request(WorkflowRequest.RequestType.REPAIR)
        self._approve_and_process(repair_id)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.REPAIRING)

        scrap_id = self._create_request(WorkflowRequest.RequestType.SCRAP)
        self._approve_and_process(scrap_id)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.status, Asset.Status.SCRAPPED)
        self.assertEqual(InventoryTransaction.objects.filter(asset=self.asset).count(), 2)

    def test_designated_approver_blocks_other_warehouse_users(self):
        designated = User.objects.create_user(username='designated', password='password123', is_staff=True)
        designated.groups.add(Group.objects.get(name='warehouse_staff'))
        request_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.BORROW,
            [{'asset': str(self.asset.id), 'quantity': 1}],
            designated_approver=str(designated.id),
        )
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        self.client.force_authenticate(self.operator)
        response = self.client.post(f'/api/workflow/requests/{request_id}/approve/', {}, format='json')
        self.assertEqual(response.status_code, 404)

        self.client.force_authenticate(designated)
        response = self.client.post(f'/api/workflow/requests/{request_id}/approve/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(ApprovalTask.objects.get(request_id=request_id).approver_id, designated.id)

    def test_applicant_can_withdraw_pending_request_and_hide_it_afterward(self):
        request_id = self._create_request()
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')

        response = self.client.post(f'/api/workflow/requests/{request_id}/withdraw/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['status'], WorkflowRequest.Status.CANCELED)
        self.assertEqual(ApprovalTask.objects.get(request_id=request_id).status, ApprovalTask.Status.CANCELED)

        response = self.client.post(f'/api/workflow/requests/{request_id}/hide/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(WorkflowRequest.objects.get(pk=request_id).is_hidden_by_applicant)

        response = self.client.get('/api/workflow/requests/')
        self.assertEqual(response.data['results'], [])

        self.client.force_authenticate(user=None)
        self.client.login(username='employee', password='password123')
        response = self.client.get(f'/workflow/{request_id}/')
        self.assertEqual(response.status_code, 404)

    def test_asset_is_reserved_on_submit_and_released_on_withdraw(self):
        request_id = self._create_request()
        response = self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        reservation = InventoryReservation.objects.get(request_line__request_id=request_id)
        self.assertEqual(reservation.status, InventoryReservation.Status.ACTIVE)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.availability_state, Asset.AvailabilityState.RESERVED)

        second_request_id = self._create_request()
        response = self.client.post(f'/api/workflow/requests/{second_request_id}/submit/', {}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('已被其他申请预约', response.data['detail'])

        self.client.post(f'/api/workflow/requests/{request_id}/withdraw/', {}, format='json')
        reservation.refresh_from_db()
        self.assertEqual(reservation.status, InventoryReservation.Status.RELEASED)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.availability_state, Asset.AvailabilityState.AVAILABLE)
        response = self.client.post(f'/api/workflow/requests/{second_request_id}/submit/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)

    def test_active_reservation_migration_backfills_asset_availability(self):
        request_id = self._create_request()
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        Asset.objects.filter(pk=self.asset.pk).update(
            availability_state=Asset.AvailabilityState.AVAILABLE,
        )

        migration = import_module('apps.workflow.migrations.0019_backfill_active_asset_reservations')
        migration.backfill_active_asset_reservations(apps, None)

        self.asset.refresh_from_db()
        self.assertEqual(self.asset.availability_state, Asset.AvailabilityState.RESERVED)

    def test_quick_process_cannot_take_asset_reserved_by_another_request(self):
        request_id = self._create_request()
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        quick_request = WorkflowRequest.objects.create(
            request_type=WorkflowRequest.RequestType.ISSUE,
            applicant=self.operator,
            reason='电话代办',
            recipient_name='销售人员',
            recipient_phone='13800138000',
            usage_location='销售部',
            is_quick_process=True,
            contact_source='电话',
        )
        quick_request.lines.create(asset=self.asset, quantity=1)

        with self.assertRaisesMessage(WorkflowError, '已被申请'):
            quick_process_request(quick_request, self.operator)

    def test_stock_reservation_uses_available_quantity_and_creates_loan(self):
        first_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.BORROW,
            [{'stock_item': str(self.stock_item.id), 'quantity': 4}],
        )
        response = self.client.post(f'/api/workflow/requests/{first_id}/submit/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(
            InventoryReservation.objects.get(request_line__request_id=first_id).quantity,
            4,
        )
        second_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.ISSUE,
            [{'stock_item': str(self.stock_item.id), 'quantity': 2}],
        )
        response = self.client.post(f'/api/workflow/requests/{second_id}/submit/', {}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('可申请 1', response.data['detail'])

        self.client.force_authenticate(self.operator)
        self.client.post(f'/api/workflow/requests/{first_id}/approve/', {}, format='json')
        response = self.client.post(f'/api/workflow/requests/{first_id}/process/', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        reservation = InventoryReservation.objects.get(request_line__request_id=first_id)
        self.assertEqual(reservation.status, InventoryReservation.Status.CONSUMED)
        loan = StockItemLoan.objects.get(source_line__request_id=first_id)
        self.assertEqual(loan.outstanding_quantity, 4)

    def test_expired_pending_request_is_closed_but_approved_reservation_is_preserved(self):
        pending_id = self._create_request()
        self.client.post(f'/api/workflow/requests/{pending_id}/submit/', {}, format='json')
        pending_reservation = InventoryReservation.objects.get(request_line__request_id=pending_id)
        self.assertIsNotNone(pending_reservation.expires_at)
        InventoryReservation.objects.filter(pk=pending_reservation.pk).update(
            expires_at=timezone.now() - timedelta(minutes=1),
        )

        self.assertEqual(expire_pending_reservations(), 1)
        self.assertEqual(
            WorkflowRequest.objects.get(pk=pending_id).status,
            WorkflowRequest.Status.CLOSED,
        )
        pending_reservation.refresh_from_db()
        self.assertEqual(pending_reservation.status, InventoryReservation.Status.RELEASED)
        self.assertIn('超过有效期', pending_reservation.release_reason)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.availability_state, Asset.AvailabilityState.AVAILABLE)

        second_asset = Asset.objects.create(
            name='测试载荷 02',
            item_type_id=self.asset.item_type_id,
            warehouse=self.warehouse,
            department=self.department,
        )
        approved_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.BORROW,
            [{'asset': str(second_asset.id), 'quantity': 1}],
        )
        self.client.post(f'/api/workflow/requests/{approved_id}/submit/', {}, format='json')
        self.client.force_authenticate(self.operator)
        self.client.post(f'/api/workflow/requests/{approved_id}/approve/', {}, format='json')
        approved_reservation = InventoryReservation.objects.get(request_line__request_id=approved_id)
        self.assertIsNone(approved_reservation.expires_at)
        self.assertEqual(expire_pending_reservations(timezone.now() + timedelta(days=30)), 0)
        self.assertEqual(
            WorkflowRequest.objects.get(pk=approved_id).status,
            WorkflowRequest.Status.WAITING_WAREHOUSE,
        )

    def test_manual_close_records_release_actor_and_reason(self):
        request_id = self._create_request()
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        request_obj = WorkflowRequest.objects.get(pk=request_id)

        close_request_and_release_reservations(request_obj, self.operator, '项目取消')

        reservation = InventoryReservation.objects.get(request_line__request_id=request_id)
        self.assertEqual(reservation.released_by, self.operator)
        self.assertEqual(reservation.release_reason, '项目取消')
        self.assertEqual(ApprovalTask.objects.get(request_id=request_id).status, ApprovalTask.Status.CANCELED)

    def test_loan_extension_approval_updates_request_due_date(self):
        request_id = self._create_request()
        self._approve_and_process(request_id)
        request_obj = WorkflowRequest.objects.get(pk=request_id)
        requested_date = request_obj.expected_return_date + timedelta(days=5)

        extension = create_loan_extension_request(
            request_obj,
            self.employee,
            requested_date,
            '现场测试尚未结束',
        )
        review_loan_extension(extension, self.operator, True, '同意延期')

        extension.refresh_from_db()
        request_obj.refresh_from_db()
        self.assertEqual(extension.status, LoanExtensionRequest.Status.APPROVED)
        self.assertEqual(request_obj.expected_return_date, requested_date)

    def test_quick_stock_issue_creates_done_request_and_cost_snapshot(self):
        self.stock_item.unit_cost = 12.50
        self.stock_item.save(update_fields=['unit_cost', 'updated_at'])
        request_obj = WorkflowRequest.objects.create(
            request_type=WorkflowRequest.RequestType.ISSUE,
            applicant=self.operator,
            reason='现场紧急补发', recipient_name='现场工程师', recipient_phone='13800138000',
            usage_location='客户现场', project_code='P-001', work_order_no='WO-001',
            is_quick_process=True, contact_source='电话',
        )
        request_obj.lines.create(stock_item=self.stock_item, quantity=2)
        quick_process_request(request_obj, self.operator)
        request_obj.refresh_from_db()
        self.stock_item.refresh_from_db()
        transaction = InventoryTransaction.objects.get(request=request_obj)
        self.assertEqual(request_obj.status, WorkflowRequest.Status.DONE)
        self.assertEqual(self.stock_item.quantity, 3)
        self.assertEqual(transaction.cost_amount, 25)
        self.assertEqual(transaction.project_code, 'P-001')

    def test_api_cannot_set_request_status_or_applicant_directly(self):
        self.client.force_authenticate(self.employee)
        response = self.client.post('/api/workflow/requests/', {
            'request_type': WorkflowRequest.RequestType.BORROW,
            'reason': '状态字段必须由流程服务更新',
            'status': WorkflowRequest.Status.DONE,
            'applicant': str(self.operator.id),
            'lines': [{'asset': str(self.asset.id), 'quantity': 1}],
        }, format='json')

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data['status'], WorkflowRequest.Status.DRAFT)
        self.assertEqual(response.data['applicant'], self.employee.id)

    def test_outbound_request_requires_traceability_before_submission(self):
        self.client.force_authenticate(self.employee)
        response = self.client.post('/api/workflow/requests/', {
            'request_type': WorkflowRequest.RequestType.BORROW,
            'reason': '缺少追溯信息测试',
            'lines': [{'asset': str(self.asset.id), 'quantity': 1}],
        }, format='json')
        self.assertEqual(response.status_code, 201, response.data)
        request_id = response.data['id']
        response = self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn('使用地点', response.data['detail'])

    def test_operator_cannot_list_other_users_designated_task_details(self):
        designated = User.objects.create_user(username='designated-logs', password='password123', is_staff=True)
        designated.groups.add(Group.objects.get(name='warehouse_staff'))
        request_id = self._create_request_with_lines(
            WorkflowRequest.RequestType.BORROW,
            [{'asset': str(self.asset.id), 'quantity': 1}],
            designated_approver=str(designated.id),
        )
        self.client.post(f'/api/workflow/requests/{request_id}/submit/', {}, format='json')
        self.client.force_authenticate(self.operator)

        response = self.client.get('/api/workflow/approval-logs/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['results'], [])


class PurchaseOrderWorkflowTests(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(username='purchase-manager', password='password123')
        self.manager.groups.add(Group.objects.get_or_create(name='warehouse_staff')[0])
        self.executor = User.objects.create_user(username='purchase-executor', password='password123')
        self.executor.groups.add(Group.objects.get_or_create(name='warehouse_outbound')[0])
        self.approver = User.objects.create_user(username='purchase-approver', password='password123')
        self.approver.groups.add(Group.objects.get_or_create(name='warehouse_approval')[0])
        self.employee = User.objects.create_user(username='purchase-employee', password='password123')
        self.warehouse = Warehouse.objects.create(name='采购测试仓', code='PURCHASE-WH')
        self.other_warehouse = Warehouse.objects.create(name='采购其他仓', code='PURCHASE-OTHER')
        self.location = Location.objects.create(warehouse=self.warehouse, name='采购测试库位', code='PUR-A01')
        self.other_location = Location.objects.create(warehouse=self.other_warehouse, name='采购其他库位', code='PUR-B01')
        category = Category.objects.create(name='采购测试分类', code='PURCHASE-CAT')
        self.asset_type = ItemType.objects.create(name='采购测试设备', code='PURCHASE-ASSET', category=category, is_serialized=True, unit='台')
        self.stock_type = ItemType.objects.create(name='采购测试配件类型', code='PURCHASE-STOCK-TYPE', category=category, is_serialized=False, unit='个')
        self.stock = StockItem.objects.create(
            name='采购测试配件', code='PURCHASE-STOCK', item_type=self.stock_type,
            quantity=2, unit='个', unit_cost='15.00', warehouse=self.warehouse, location=self.location,
        )

    def _approved_request(self, *, source, quantity=1):
        purchase_request = PurchaseRequest.objects.create(
            applicant=self.employee,
            reason='采购闭环专项测试',
            status=PurchaseRequest.Status.PENDING_APPROVAL,
        )
        if source == 'stock':
            purchase_request.lines.create(
                stock_item=self.stock, name_snapshot=self.stock.name,
                code_snapshot=self.stock.code, management_mode='耗材配件', quantity=quantity,
                unit='个', reference_unit_cost='15.00', supplier_snapshot='测试供应商',
            )
        else:
            purchase_request.lines.create(
                item_type=self.asset_type, name_snapshot=self.asset_type.name,
                code_snapshot=self.asset_type.code, management_mode='单件资产', quantity=quantity,
                unit='台', reference_unit_cost='100.00', supplier_snapshot='测试设备供应商',
            )
        approve_purchase_request(purchase_request, self.approver, '同意采购')
        return purchase_request

    def test_only_approved_request_can_create_order_and_employee_is_rejected(self):
        draft = PurchaseRequest.objects.create(applicant=self.employee, reason='尚未审批')
        with self.assertRaisesMessage(WorkflowError, '只有具备录入与基础资料权限'):
            create_purchase_order(draft, self.employee)
        draft.status = PurchaseRequest.Status.APPROVED
        draft.save(update_fields=['status', 'updated_at'])
        with self.assertRaisesMessage(WorkflowError, '采购申请没有有效明细'):
            create_purchase_order(draft, self.manager)

    def test_stock_partial_receipt_updates_quantity_order_and_transaction(self):
        request_obj = self._approved_request(source='stock', quantity=5)
        order = create_purchase_order(request_obj, self.manager)
        self.assertEqual(order.status, PurchaseOrder.Status.DRAFT)
        place_purchase_order(order, self.manager, date.today() + timedelta(days=7), '测试下达')
        receipt = receive_purchase_order(
            order, self.executor, {str(order.lines.get().id): 2}, self.warehouse, self.location,
        )
        self.stock.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(receipt.inventory_request.request_type, WorkflowRequest.RequestType.PURCHASE_RECEIPT)
        self.assertEqual(receipt.inventory_request.status, WorkflowRequest.Status.DONE)
        self.assertEqual(receipt.inventory_request.lines.count(), 1)
        self.assertEqual(order.status, PurchaseOrder.Status.PARTIAL_RECEIVED)
        self.assertEqual(self.stock.quantity, 4)
        self.assertEqual(InventoryTransaction.objects.filter(request=receipt.inventory_request).count(), 1)
        self.assertEqual(PurchaseReceipt.objects.filter(order=order).count(), 1)

    def test_asset_receipt_creates_each_asset_and_finishes_order(self):
        request_obj = self._approved_request(source='asset', quantity=2)
        order = create_purchase_order(request_obj, self.manager)
        place_purchase_order(order, self.manager)
        receipt = receive_purchase_order(
            order, self.executor, {str(order.lines.get().id): 2}, self.warehouse, self.location,
        )
        order.refresh_from_db()
        assets = Asset.objects.filter(purchase_order_no=order.order_no)
        self.assertEqual(order.status, PurchaseOrder.Status.RECEIVED)
        self.assertEqual(assets.count(), 2)
        self.assertEqual(assets.filter(status=Asset.Status.IN_STOCK, warehouse=self.warehouse, location=self.location).count(), 2)
        self.assertEqual(InventoryTransaction.objects.filter(request=receipt.inventory_request, asset__isnull=False).count(), 2)
        self.assertEqual(receipt.inventory_request.lines.count(), 2)
        self.assertEqual(len(receipt.lines.get().created_asset_ids), 2)

    def test_receipt_rejects_wrong_location_and_excess_quantity_without_inventory_write(self):
        request_obj = self._approved_request(source='stock', quantity=2)
        order = create_purchase_order(request_obj, self.manager)
        place_purchase_order(order, self.manager)
        line_id = str(order.lines.get().id)
        with self.assertRaisesMessage(WorkflowError, '库位不属于所选仓库'):
            receive_purchase_order(order, self.executor, {line_id: 1}, self.warehouse, self.other_location)
        with self.assertRaisesMessage(WorkflowError, '最多还能到货 2'):
            receive_purchase_order(order, self.executor, {line_id: 3}, self.warehouse, self.location)
        self.stock.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(self.stock.quantity, 2)
        self.assertEqual(order.lines.get().received_quantity, 0)
        self.assertEqual(PurchaseReceipt.objects.count(), 0)
        self.assertEqual(InventoryTransaction.objects.count(), 0)

    def test_purchase_order_pages_render_for_inventory_manager(self):
        request_obj = self._approved_request(source='stock', quantity=1)
        order = create_purchase_order(request_obj, self.manager)
        self.client.force_login(self.manager)
        self.assertEqual(self.client.get('/warehouse/purchase-orders/').status_code, 200)
        response = self.client.get(f'/warehouse/purchase-orders/{order.id}/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, order.order_no)


class AssetProfileCompletionTests(TestCase):
    """采购到货资产资料补录：到货登记逐件录入 + 事后批量补录。"""

    def setUp(self):
        self.manager = User.objects.create_user(username='profile-manager', password='password123')
        self.manager.groups.add(Group.objects.get_or_create(name='warehouse_staff')[0])
        self.executor = User.objects.create_user(username='profile-executor', password='password123')
        self.executor.groups.add(Group.objects.get_or_create(name='warehouse_outbound')[0])
        self.approver = User.objects.create_user(username='profile-approver', password='password123')
        self.approver.groups.add(Group.objects.get_or_create(name='warehouse_approval')[0])
        self.employee = User.objects.create_user(username='profile-employee', password='password123')
        self.warehouse = Warehouse.objects.create(name='补录测试仓', code='PROFILE-WH')
        self.location = Location.objects.create(warehouse=self.warehouse, name='补录库位', code='P-A01')
        category = Category.objects.create(name='补录分类', code='PROFILE-CAT')
        self.asset_type = ItemType.objects.create(name='补录设备', code='PROFILE-ASSET', category=category, is_serialized=True, unit='台')

    def _received_order(self, quantity=2):
        from .services import bulk_update_asset_profiles  # noqa: F401  # 确认服务可导入
        purchase_request = PurchaseRequest.objects.create(
            applicant=self.employee, reason='补录测试', status=PurchaseRequest.Status.PENDING_APPROVAL,
        )
        purchase_request.lines.create(
            item_type=self.asset_type, name_snapshot=self.asset_type.name,
            code_snapshot=self.asset_type.code, management_mode='单件资产', quantity=quantity,
            unit='台', reference_unit_cost='100.00', supplier_snapshot='补录供应商',
        )
        approve_purchase_request(purchase_request, self.approver, '同意')
        order = create_purchase_order(purchase_request, self.manager)
        place_purchase_order(order, self.manager, date.today(), '下达')
        return order

    def test_receive_with_asset_profiles(self):
        from .services import receive_purchase_order as receive
        order = self._received_order(quantity=2)
        line = order.lines.get()
        receipt = receive(
            order, self.executor, {str(line.id): 2}, self.warehouse, self.location,
            asset_profiles={str(line.id): [
                {'serial_number': 'SN-001', 'asset_code': 'ERP-001', 'manufacturer': '厂家A', 'model': 'M1'},
                {'serial_number': 'SN-002', 'asset_code': 'ERP-002'},
            ]},
        )
        assets = Asset.objects.filter(purchase_order_no=order.order_no).order_by('system_asset_no')
        self.assertEqual(assets.count(), 2)
        self.assertEqual(assets[0].serial_number, 'SN-001')
        self.assertEqual(assets[0].asset_code, 'ERP-001')
        self.assertEqual(assets[0].manufacturer, '厂家A')
        self.assertEqual(assets[1].serial_number, 'SN-002')
        self.assertEqual(assets[1].remarks, '采购到货')
        self.assertIsNotNone(receipt)

    def test_receive_rejects_duplicate_sn_in_batch(self):
        from .services import receive_purchase_order as receive
        order = self._received_order(quantity=2)
        line = order.lines.get()
        with self.assertRaisesMessage(WorkflowError, '重复'):
            receive(
                order, self.executor, {str(line.id): 2}, self.warehouse, self.location,
                asset_profiles={str(line.id): [{'serial_number': 'SN-X'}, {'serial_number': 'SN-X'}]},
            )
        self.assertEqual(Asset.objects.filter(purchase_order_no=order.order_no).count(), 0)

    def test_receive_rejects_sn_conflict_with_existing_asset(self):
        from .services import receive_purchase_order as receive
        Asset.objects.create(name='存量资产', item_type=self.asset_type, warehouse=self.warehouse, serial_number='SN-OLD')
        order = self._received_order(quantity=1)
        line = order.lines.get()
        with self.assertRaisesMessage(WorkflowError, '与台账已有资产重复'):
            receive(
                order, self.executor, {str(line.id): 1}, self.warehouse, self.location,
                asset_profiles={str(line.id): [{'serial_number': 'SN-OLD'}]},
            )

    def test_receive_rejects_excess_profiles(self):
        from .services import receive_purchase_order as receive
        order = self._received_order(quantity=1)
        line = order.lines.get()
        with self.assertRaisesMessage(WorkflowError, '超过本次到货数量'):
            receive(
                order, self.executor, {str(line.id): 1}, self.warehouse, self.location,
                asset_profiles={str(line.id): [{'serial_number': 'S1'}, {'serial_number': 'S2'}]},
            )

    def test_bulk_update_profiles(self):
        from .services import bulk_update_asset_profiles
        order = self._received_order(quantity=2)
        line = order.lines.get()
        receive_purchase_order(order, self.executor, {str(line.id): 2}, self.warehouse, self.location)
        assets = list(Asset.objects.filter(purchase_order_no=order.order_no).order_by('system_asset_no'))
        self.assertEqual(assets[0].remarks, '采购到货待补充资产信息')
        updated = bulk_update_asset_profiles(order, self.manager, [
            {'system_asset_no': assets[0].system_asset_no, 'serial_number': 'SN-101', 'asset_code': 'ERP-101', 'model': 'X1'},
            {'system_asset_no': assets[1].system_asset_no, 'serial_number': 'SN-102'},
        ])
        self.assertEqual(updated, 2)
        assets[0].refresh_from_db()
        self.assertEqual(assets[0].serial_number, 'SN-101')
        self.assertEqual(assets[0].asset_code, 'ERP-101')
        self.assertEqual(assets[0].model, 'X1')
        self.assertEqual(assets[0].remarks, '采购到货')
        # 审计与生命周期事件
        self.assertTrue(OperationAuditLog.objects.filter(action='purchase_asset_profile_updated').count() >= 2)
        from apps.inventory.models import AssetLifecycleEvent
        self.assertEqual(AssetLifecycleEvent.objects.filter(title='采购到货资料补录').count(), 2)

    def test_bulk_update_rejects_foreign_asset(self):
        from .services import bulk_update_asset_profiles
        order = self._received_order(quantity=1)
        line = order.lines.get()
        receive_purchase_order(order, self.executor, {str(line.id): 1}, self.warehouse, self.location)
        outsider = Asset.objects.create(name='外来资产', item_type=self.asset_type, warehouse=self.warehouse)
        with self.assertRaisesMessage(WorkflowError, '不属于采购订单'):
            bulk_update_asset_profiles(order, self.manager, [
                {'system_asset_no': outsider.system_asset_no, 'serial_number': 'SN-X'},
            ])

    def test_bulk_update_rejects_duplicate_and_permission(self):
        from .services import bulk_update_asset_profiles
        order = self._received_order(quantity=2)
        line = order.lines.get()
        receive_purchase_order(order, self.executor, {str(line.id): 2}, self.warehouse, self.location)
        assets = list(Asset.objects.filter(purchase_order_no=order.order_no).order_by('system_asset_no'))
        # 批次内 SN 重复
        with self.assertRaisesMessage(WorkflowError, '重复'):
            bulk_update_asset_profiles(order, self.manager, [
                {'system_asset_no': assets[0].system_asset_no, 'serial_number': 'SN-DUP'},
                {'system_asset_no': assets[1].system_asset_no, 'serial_number': 'SN-DUP'},
            ])
        # 无权限账号
        with self.assertRaisesMessage(WorkflowError, '录入与基础资料权限'):
            bulk_update_asset_profiles(order, self.employee, [
                {'system_asset_no': assets[0].system_asset_no, 'serial_number': 'SN-OK'},
            ])
        # 确认事务回滚：资产未被改动
        assets[0].refresh_from_db()
        self.assertEqual(assets[0].serial_number, '')

    def test_pending_profile_assets_query(self):
        from .services import order_pending_profile_assets
        order = self._received_order(quantity=2)
        line = order.lines.get()
        receive_purchase_order(
            order, self.executor, {str(line.id): 2}, self.warehouse, self.location,
            asset_profiles={str(line.id): [{'serial_number': 'SN-1', 'asset_code': 'E-1'}]},
        )
        pending = order_pending_profile_assets(order)
        self.assertEqual(pending.count(), 1)  # 第一件资料齐全，第二件待补录

    def test_order_detail_shows_pending_assets_and_template_download(self):
        order = self._received_order(quantity=1)
        line = order.lines.get()
        receive_purchase_order(order, self.executor, {str(line.id): 1}, self.warehouse, self.location)
        asset = Asset.objects.get(purchase_order_no=order.order_no)
        self.client.force_login(self.manager)
        response = self.client.get(f'/warehouse/purchase-orders/{order.id}/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '资产资料补录')
        self.assertContains(response, asset.system_asset_no)
        # 模板下载：用 openpyxl 读回验证预填内容
        tpl = self.client.get(f'/warehouse/purchase-orders/{order.id}/profile-template/')
        self.assertEqual(tpl.status_code, 200)
        self.assertIn('spreadsheetml', tpl['Content-Type'])
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(tpl.content))
        sheet = wb.worksheets[0]
        self.assertEqual(sheet.cell(row=1, column=1).value, '系统资产编号')
        self.assertEqual(sheet.cell(row=2, column=1).value, asset.system_asset_no)

    def test_order_detail_bulk_upload_completes_profiles(self):
        from io import BytesIO
        from openpyxl import Workbook
        order = self._received_order(quantity=1)
        line = order.lines.get()
        receive_purchase_order(order, self.executor, {str(line.id): 1}, self.warehouse, self.location)
        asset = Asset.objects.get(purchase_order_no=order.order_no)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['系统资产编号', '生产厂家 SN', '资产编码', '生产厂家条码', '生产厂家', '型号'])
        sheet.append([asset.system_asset_no, 'SN-UP-1', 'ERP-UP-1', '', '厂家B', 'M2'])
        buf = BytesIO()
        workbook.save(buf)
        buf.seek(0)
        self.client.force_login(self.manager)
        from django.core.files.uploadedfile import SimpleUploadedFile
        upload = SimpleUploadedFile('profiles.xlsx', buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response = self.client.post(
            f'/warehouse/purchase-orders/{order.id}/',
            {'action': 'complete_profiles', 'profile_file': upload},
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        asset.refresh_from_db()
        msgs = [str(m) for m in list(response.context.get('messages', []))] if response.context else []
        self.assertEqual(asset.serial_number, 'SN-UP-1', f'资产未更新，页面消息: {msgs}')
        self.assertEqual(asset.asset_code, 'ERP-UP-1')
        self.assertEqual(asset.manufacturer, '厂家B')

    def test_order_detail_receive_with_inline_profiles(self):
        order = self._received_order(quantity=2)
        line = order.lines.get()
        place_kwargs = {'action': 'receive', 'warehouse_id': str(self.warehouse.id), 'location_id': str(self.location.id),
                        'received_date': date.today().isoformat(), f'received_{line.id}': '2',
                        f'profile_{line.id}_1_serial_number': 'SN-W1', f'profile_{line.id}_1_asset_code': 'ERP-W1',
                        f'profile_{line.id}_2_serial_number': 'SN-W2'}
        self.client.force_login(self.executor)
        response = self.client.post(f'/warehouse/purchase-orders/{order.id}/', place_kwargs)
        self.assertEqual(response.status_code, 302)
        assets = Asset.objects.filter(purchase_order_no=order.order_no).order_by('system_asset_no')
        self.assertEqual(assets.count(), 2)
        self.assertEqual(assets[0].serial_number, 'SN-W1')
        self.assertEqual(assets[0].asset_code, 'ERP-W1')
        self.assertEqual(assets[1].serial_number, 'SN-W2')

    def test_employee_cannot_download_template_or_upload(self):
        order = self._received_order(quantity=1)
        line = order.lines.get()
        receive_purchase_order(order, self.executor, {str(line.id): 1}, self.warehouse, self.location)
        self.client.force_login(self.employee)
        tpl = self.client.get(f'/warehouse/purchase-orders/{order.id}/profile-template/')
        self.assertIn(tpl.status_code, (302, 403))


class SupplierReportTests(TestCase):
    """供应商价格历史与对账。"""

    def setUp(self):
        self.reporter = User.objects.create_user(username='reporter', password='password123')
        self.reporter.groups.add(Group.objects.get_or_create(name='warehouse_reports')[0])
        self.employee = User.objects.create_user(username='plain', password='password123')
        self.client.force_login(self.reporter)
        category = Category.objects.create(name='报表分类', code='REPORT-CAT')
        self.asset_type = ItemType.objects.create(name='报表设备', code='REPORT-ASSET', category=category, is_serialized=True, unit='台')
        self.warehouse = Warehouse.objects.create(name='报表测试仓', code='REPORT-WH')
        self.location = Location.objects.create(warehouse=self.warehouse, name='报表测试库位', code='REPORT-A01')

    def _make_order(self, order_date, supplier, item_name, item_code, unit_cost, ordered, received):
        from .models import PurchaseRequestLine
        purchase_request = PurchaseRequest.objects.create(
            applicant=self.employee, reason='报表测试', status=PurchaseRequest.Status.APPROVED,
        )
        request_line = PurchaseRequestLine.objects.create(
            purchase_request=purchase_request,
            item_type=self.asset_type,
            name_snapshot=item_name, code_snapshot=item_code,
            management_mode='单件资产', quantity=ordered, unit='台',
            reference_unit_cost=unit_cost, supplier_snapshot=supplier,
        )
        order = PurchaseOrder.objects.create(
            purchase_request=purchase_request,
            order_date=order_date,
            status=PurchaseOrder.Status.RECEIVED if received == ordered else PurchaseOrder.Status.PARTIAL_RECEIVED,
        )
        line = order.lines.create(
            purchase_request_line=request_line,
            item_type=self.asset_type,
            name_snapshot=item_name,
            code_snapshot=item_code,
            management_mode='单件资产',
            ordered_quantity=ordered,
            received_quantity=received,
            unit='台',
            unit_cost=unit_cost,
            supplier_snapshot=supplier,
        )
        return order, line

    def test_price_history_groups_by_supplier_and_item(self):
        from decimal import Decimal
        self._make_order(date(2026, 8, 1), '供应商A', '电机', 'MTR-01', Decimal('100.00'), 2, 2)
        self._make_order(date(2026, 9, 1), '供应商A', '电机', 'MTR-01', Decimal('95.00'), 3, 3)
        self._make_order(date(2026, 9, 5), '供应商B', '电机', 'MTR-01', Decimal('110.00'), 1, 1)
        response = self.client.get('/reports/supplier-prices/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '供应商A')
        self.assertContains(response, '95.00')  # 最新价
        # 筛选供应商
        response = self.client.get('/reports/supplier-prices/', {'supplier': '供应商A'})
        self.assertNotContains(response, '110.00')

    def test_price_history_export_csv(self):
        from decimal import Decimal
        self._make_order(date(2026, 8, 1), '供应商A', '电机', 'MTR-01', Decimal('100.00'), 2, 2)
        response = self.client.get('/reports/supplier-prices.csv')
        self.assertEqual(response.status_code, 200)
        self.assertIn('text/csv', response['Content-Type'])
        content = response.content.decode('utf-8-sig')
        self.assertIn('供应商A', content)
        self.assertIn('100.00', content)

    def test_reconciliation_summary(self):
        from decimal import Decimal
        self._make_order(date(2026, 8, 1), '供应商A', '电机', 'MTR-01', Decimal('100.00'), 2, 2)
        self._make_order(date(2026, 8, 10), '供应商A', '减速器', 'GB-01', Decimal('50.00'), 4, 2)
        response = self.client.get('/reports/supplier-reconciliation/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '供应商A')
        self.assertContains(response, '400.00')  # 订购金额 100*2 + 50*4
        self.assertContains(response, '300.00')  # 到货金额 100*2 + 50*2

    def test_reconciliation_export_csv(self):
        from decimal import Decimal
        self._make_order(date(2026, 8, 1), '供应商A', '电机', 'MTR-01', Decimal('100.00'), 2, 2)
        response = self.client.get('/reports/supplier-reconciliation.csv')
        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8-sig')
        self.assertIn('供应商A', content)
        self.assertIn('200.00', content)

    def test_employee_cannot_access_supplier_reports(self):
        self.client.force_login(self.employee)
        for url in ('/reports/supplier-prices/', '/reports/supplier-reconciliation/'):
            response = self.client.get(url)
            self.assertIn(response.status_code, (302, 403), f'{url} 应拒绝普通员工')

    def test_reports_exclude_draft_and_canceled_orders(self):
        from decimal import Decimal
        canceled, _ = self._make_order(date(2026, 8, 1), '供应商A', '取消物品', 'CANCEL-01', Decimal('999.00'), 1, 0)
        canceled.status = PurchaseOrder.Status.CANCELED
        canceled.save(update_fields=['status', 'updated_at'])
        draft, _ = self._make_order(date(2026, 8, 2), '供应商A', '草稿物品', 'DRAFT-01', Decimal('888.00'), 1, 0)
        draft.status = PurchaseOrder.Status.DRAFT
        draft.save(update_fields=['status', 'updated_at'])
        response = self.client.get('/reports/supplier-prices/')
        self.assertNotContains(response, '999.00')
        self.assertNotContains(response, '888.00')
        response = self.client.get('/reports/supplier-reconciliation/')
        self.assertNotContains(response, '取消物品')
        self.assertNotContains(response, '草稿物品')

    def test_reconciliation_can_filter_by_receipt_date(self):
        from decimal import Decimal
        order, line = self._make_order(date(2026, 8, 1), '供应商A', '跨月到货物品', 'CROSS-01', Decimal('100.00'), 2, 2)
        receipt = PurchaseReceipt.objects.create(
            order=order,
            received_by=self.employee,
            received_date=date(2026, 9, 3),
        )
        PurchaseReceiptLine.objects.create(
            receipt=receipt,
            order_line=line,
            received_quantity=2,
            destination_warehouse=self.warehouse,
            destination_location=self.location,
        )
        response = self.client.get('/reports/supplier-reconciliation/', {
            'date_basis': 'receipt', 'start_date': '2026-09-01', 'end_date': '2026-09-30',
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, '跨月到货物品')
        self.assertContains(response, '200.00')
